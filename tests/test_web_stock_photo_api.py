import io
import sqlite3
import threading
from http.server import ThreadingHTTPServer

import requests
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.test_vendor_management_center import _setup_tenant
import Web_Arayuzu.Main as web_main


def test_stock_photo_route_requires_login_and_reads_only_active_tenant(monkeypatch, tmp_path):
    _actor, tenant_path = _setup_tenant(monkeypatch, tmp_path)
    photo_path = tmp_path / "stock-photo.png"
    photo_path.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
        b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
        b"\x1f\x15\xc4\x89\x00\x00\x00\x0bIDATx\xda\x63\xf8\x0f"
        b"\x00\x01\x01\x01\x00\x18\xdd\x8d\xb1\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    with sqlite3.connect(tenant_path) as connection:
        connection.execute(
            "CREATE TABLE IF NOT EXISTS parts (id INTEGER PRIMARY KEY, photo_path TEXT, is_deleted INTEGER DEFAULT 0)"
        )
        columns = {row[1] for row in connection.execute("PRAGMA table_info(parts)")}
        for column, declaration in (
            ("photo_path", "TEXT"), ("is_deleted", "INTEGER DEFAULT 0"),
            ("stock", "REAL DEFAULT 0"), ("purchase_price", "REAL DEFAULT 0"),
            ("currency", "TEXT DEFAULT 'TRY'"), ("name", "TEXT"),
        ):
            if column not in columns:
                connection.execute(f"ALTER TABLE parts ADD COLUMN {column} {declaration}")
        connection.execute(
            "INSERT INTO parts(id,photo_path,is_deleted,stock,purchase_price,currency,name) VALUES(?,?,?,?,?,?,?)",
            (91001, str(photo_path), 0, 4, 3, "TRY", "Test Stock"),
        )
        connection.execute(
            "INSERT INTO parts(id,photo_path,is_deleted,stock,purchase_price,price,currency,name) "
            "VALUES(91002,'',0,2,10,15,'USD','Dollar Stock')"
        )
        connection.execute(
            "INSERT INTO exchange_rates(currency,buying_rate,selling_rate,effective_date,created_at) "
            "VALUES('USD',32,32.5,'2026-10-14','2026-10-14 12:00:00')"
        )

    server = ThreadingHTTPServer(("127.0.0.1", 0), web_main.AYECRequestHandler)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base_url = f"http://127.0.0.1:{server.server_address[1]}"
    endpoint = f"{base_url}/api/desktop/stock/photo/91001"
    try:
        anonymous = requests.Session()
        anonymous.headers["Origin"] = base_url
        assert anonymous.get(endpoint).status_code == 401
        assert anonymous.get(f"{base_url}/api/desktop/stock/export/xlsx").status_code == 401
        assert anonymous.get(f"{base_url}/api/desktop/stock/export/pdf").status_code == 401
        assert anonymous.post(
            f"{base_url}/api/desktop/stock/delete", json={"id": 91001, "password": "operator-test-pass"}
        ).status_code == 401

        session = requests.Session()
        session.headers["Origin"] = base_url
        login = session.post(
            f"{base_url}/api/auth/login",
            json={"identifier": web_main.VENDOR_EMAIL, "password": "operator-test-pass", "tenant_id": "tenant-1"},
        )
        assert login.status_code == 200
        response = session.get(endpoint)
        assert response.status_code == 200
        assert response.headers["Content-Type"] == "image/png"
        assert response.content == photo_path.read_bytes()
        assert session.get(f"{base_url}/api/desktop/stock/photo/99999").status_code == 404
        workbook_response = session.get(f"{base_url}/api/desktop/stock/export/xlsx")
        assert workbook_response.status_code == 200
        assert "spreadsheetml.sheet" in workbook_response.headers["Content-Type"]
        workbook = load_workbook(io.BytesIO(workbook_response.content), read_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        assert sheet.title == "Stok"
        assert "photo_path" not in headers
        assert sheet.max_row >= 2
        assert "D\u00f6viz Kuru" in headers
        assert "Al\u0131\u015f Fiyat\u0131 (TRY)" in headers
        assert "Sat\u0131\u015f Fiyat\u0131 (TRY)" in headers
        assert "Stok De\u011feri (TRY)" in headers
        exported = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
        dollar_row = next(row for row in exported if row.get("\u00dcr\u00fcn Ad\u0131") == "Dollar Stock")
        assert dollar_row["D\u00f6viz Kuru"] == 32.5
        assert dollar_row["Al\u0131\u015f Fiyat\u0131 (TRY)"] == 325
        assert dollar_row["Sat\u0131\u015f Fiyat\u0131 (TRY)"] == 487.5
        assert dollar_row["Stok De\u011feri (TRY)"] == 650
        workbook.close()
        pdf_response = session.get(f"{base_url}/api/desktop/stock/export/pdf")
        assert pdf_response.status_code == 200
        assert pdf_response.headers["Content-Type"] == "application/pdf"
        assert pdf_response.content.startswith(b"%PDF-")
        pdf = PdfReader(io.BytesIO(pdf_response.content))
        pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        assert "AYEC Pro - Stok Listesi" in pdf_text
        assert "Test Stock" in pdf_text
        assert "Dollar Stock" in pdf_text
        delete_url = f"{base_url}/api/desktop/stock/delete"
        rejected = session.post(delete_url, json={"id": 91001, "password": "wrong-pass"})
        assert rejected.status_code == 403
        with sqlite3.connect(tenant_path) as connection:
            assert connection.execute("SELECT is_deleted FROM parts WHERE id=91001").fetchone()[0] == 0
        deleted = session.post(delete_url, json={"id": 91001, "password": "operator-test-pass"})
        assert deleted.status_code == 200
        assert deleted.json() == {"ok": True, "id": 91001}
        with sqlite3.connect(tenant_path) as connection:
            assert connection.execute("SELECT is_deleted FROM parts WHERE id=91001").fetchone()[0] == 1
            assert connection.execute(
                "SELECT COUNT(*) FROM stock_movements WHERE part_id=91001 AND movement_type=?",
                ("\u00c7\u0131k\u0131\u015f (Silindi)",),
            ).fetchone()[0] == 1
            assert connection.execute(
                "SELECT COUNT(*) FROM accounting WHERE category=? AND description LIKE ?",
                ("Stok \u0130ptali / Silinme", "%Test Stock%"),
            ).fetchone()[0] == 1
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)
