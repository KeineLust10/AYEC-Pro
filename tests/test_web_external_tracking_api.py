import base64
import io
import sqlite3
import threading
from http.server import ThreadingHTTPServer

import requests
from openpyxl import Workbook, load_workbook

from tests.test_vendor_management_center import _setup_tenant
import Web_Arayuzu.Main as web_main


def test_external_tracking_crud_uses_active_tenant(monkeypatch, tmp_path):
    _actor, tenant_path = _setup_tenant(monkeypatch, tmp_path)
    server = ThreadingHTTPServer(("127.0.0.1", 0), web_main.AYECRequestHandler)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base_url = f"http://127.0.0.1:{server.server_address[1]}"
    table_url = f"{base_url}/api/desktop/table/external_warranty_tracking"
    try:
        anonymous = requests.Session()
        anonymous.headers["Origin"] = base_url
        assert anonymous.get(table_url).status_code == 401
        assert anonymous.get(f"{base_url}/api/desktop/documents/center").status_code == 401

        session = requests.Session()
        session.headers["Origin"] = base_url
        login = session.post(
            f"{base_url}/api/auth/login",
            json={"identifier": web_main.VENDOR_EMAIL, "password": "operator-test-pass", "tenant_id": "tenant-1"},
        )
        assert login.status_code == 200

        payload = {
            "internal_tracking_no": "M.NO 1025",
            "customer_name": "Example Customer",
            "product_name": "Laptop board",
            "external_service_name": "Service One",
            "external_service_no": "SV-205",
            "sent_date": "2026-10-13",
            "outbound_cargo_no": "OUT-01",
            "inbound_cargo_no": "IN-02",
            "status": "Kargoya Verildi",
            "notes": "Check connector",
        }
        added = session.post(table_url, json=payload)
        assert added.status_code == 200
        record_id = added.json()["id"]
        found = session.get(f"{table_url}?q=SV-205").json()
        assert [row["id"] for row in found["rows"]] == [record_id]

        changed = session.post(
            table_url,
            json={"id": record_id, "_action": "update", "status": "Onarımda", "inbound_cargo_no": "IN-03"},
        )
        assert changed.status_code == 200
        updated = session.get(f"{table_url}?q=IN-03").json()["rows"][0]
        assert updated["status"] == "Onarımda"

        removed = session.post(table_url, json={"id": record_id, "_action": "delete"})
        assert removed.status_code == 200
        assert session.get(table_url).json()["rows"] == []
        with sqlite3.connect(tenant_path) as conn:
            assert conn.execute("SELECT COUNT(*) FROM external_warranty_tracking").fetchone()[0] == 0

        photos_url = f"{base_url}/api/desktop/external-tracking/photos"
        uploaded = session.post(
            photos_url,
            json={
                "tracking_no": "M.NO 1025",
                "label": "sample.png",
                "stage": "Giri\u015f",
                "data": "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p4sAAAAASUVORK5CYII=",
            },
        )
        assert uploaded.status_code == 201
        photo_id = uploaded.json()["id"]
        with sqlite3.connect(tenant_path) as conn:
            assert conn.execute("SELECT stage FROM photos WHERE id=?", (photo_id,)).fetchone()[0] == "Giri\u015f"
        listed = session.get(f"{photos_url}?tracking_no=M.NO%201025").json()["rows"]
        assert [photo["id"] for photo in listed] == [photo_id]
        documents = session.get(f"{base_url}/api/desktop/documents/center")
        assert documents.status_code == 200
        assert any(row["key"] == f"photo:{photo_id}" and row["available"] for row in documents.json()["rows"])
        assert all("file_path" not in row and "photo_path" not in row for row in documents.json()["rows"])
        image_response = session.get(f"{base_url}/api/desktop/external-tracking/photo/{photo_id}")
        assert image_response.status_code == 200
        assert image_response.headers["Content-Type"] == "image/png"
        assert image_response.content.startswith(b"\x89PNG\r\n\x1a\n")
        assert session.post(photos_url, json={"_action": "delete", "id": photo_id}).status_code == 200
        assert session.get(f"{base_url}/api/desktop/external-tracking/photo/{photo_id}").status_code == 404

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Emanet No", "Musteri", "Urun", "Servis", "Servis No", "Tarih", "Gidis Kargo", "Donus Kargo", "Durum", "Notlar"])
        sheet.append(["IMP-1", "Import Buyer", "Router", "Service Two", "SV-9", "2026-10-12", "OUT-9", "", "", ""])
        sheet.cell(row=2, column=10).value = "=1+1"
        sheet.cell(row=2, column=10).data_type = "s"
        buffer = io.BytesIO()
        workbook.save(buffer)
        imported = session.post(
            f"{base_url}/api/desktop/external-tracking/import",
            json={"filename": "tracking.xlsx", "data": base64.b64encode(buffer.getvalue()).decode("ascii")},
        )
        assert imported.status_code == 200
        assert imported.json() == {"success": 1, "failures": 0}
        imported_row = session.get(f"{table_url}?q=IMP-1").json()["rows"][0]
        assert imported_row["status"] == "Haz\u0131rlan\u0131yor"
        assert imported_row["notes"] == "=1+1"

        exported = session.get(f"{base_url}/api/desktop/external-tracking/export?ids={imported_row['id']}")
        assert exported.status_code == 200
        assert exported.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        report = load_workbook(io.BytesIO(exported.content), data_only=False)
        report_row = list(report.active.iter_rows(min_row=2, values_only=False))[0]
        assert report_row[1].value == "IMP-1"
        assert report_row[10].value == "=1+1"
        assert report_row[10].data_type == "s"
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)


def test_external_tracking_context_menu_matches_desktop_actions():
    source = (web_main.ROOT / "web" / "external-tracking-reference.js").read_text(encoding="utf-8")
    assert 'data-external-context="edit"' in source
    assert 'data-external-context="delete"' in source
    assert 'data-external-context="gallery"' in source
    assert 'data-external-context="export"' in source
    assert 'const id = String(tr.dataset.externalRow)' in source
    assert 'menu.dataset.targetId = id' in source
    assert 'openGallery(row)' in source
    assert 'exportWorkbook().catch' in source
