import io
import threading
from http.server import ThreadingHTTPServer

import requests
from openpyxl import load_workbook

from tests.test_vendor_management_center import _setup_tenant
import Web_Arayuzu.Main as web_main


def test_service_list_exports_are_authenticated_and_produce_desktop_columns(monkeypatch, tmp_path):
    _actor, _tenant_path = _setup_tenant(monkeypatch, tmp_path)
    server = ThreadingHTTPServer(("127.0.0.1", 0), web_main.AYECRequestHandler)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base_url = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        anonymous = requests.Session()
        anonymous.headers["Origin"] = base_url
        payload = {"records": [{"tracking_no": "SRV-1", "price": 15.5}]}
        assert anonymous.post(f"{base_url}/api/desktop/service-list/export/xlsx", json=payload).status_code == 401

        session = requests.Session()
        session.headers["Origin"] = base_url
        login = session.post(
            f"{base_url}/api/auth/login",
            json={
                "identifier": web_main.VENDOR_EMAIL,
                "password": "operator-test-pass",
                "tenant_id": "tenant-1",
            },
        )
        assert login.status_code == 200
        record = {
            "tracking_no": "=1+1",
            "customer": "Test Customer",
            "company": "Test Company",
            "device_type": "Phone",
            "brand": "AYEC",
            "model": "A1",
            "entry_date": "2026-10-01",
            "exit_date": "2026-10-02",
            "price": 15.5,
            "price_display": "15,50 TRY",
            "status": "Bekliyor",
            "urgency": "Normal",
        }
        xlsx = session.post(
            f"{base_url}/api/desktop/service-list/export/xlsx",
            json={"records": [record]},
        )
        assert xlsx.status_code == 200, xlsx.text
        assert xlsx.headers["Content-Type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert "servis_listesi.xlsx" in xlsx.headers["Content-Disposition"]
        workbook = load_workbook(io.BytesIO(xlsx.content), data_only=False)
        sheet = workbook.active
        assert sheet.max_row == 2 and sheet.max_column == 11
        assert sheet.cell(1, 3).value == "\u015eirket"
        assert sheet.cell(1, 8).value == "\u00c7\u0131k\u0131\u015f Tarihi"
        assert sheet.cell(2, 1).value == "=1+1" and sheet.cell(2, 1).data_type == "s"
        assert sheet.cell(2, 9).value == 15.5
        workbook.close()

        pdf = session.post(
            f"{base_url}/api/desktop/service-list/export/pdf",
            json={"records": [record]},
        )
        assert pdf.status_code == 200
        assert pdf.headers["Content-Type"] == "application/pdf"
        assert "servis_listesi.pdf" in pdf.headers["Content-Disposition"]
        assert pdf.content.startswith(b"%PDF-")

        session.headers["Origin"] = "https://unexpected.example"
        cross_origin = session.post(
            f"{base_url}/api/desktop/service-list/export/xlsx",
            json={"records": [record]},
        )
        assert cross_origin.status_code == 403
        session.headers["Origin"] = base_url
        session.headers["X-AYEC-Fiscal-Archive"] = "invalid"
        invalid_archive = session.post(
            f"{base_url}/api/desktop/service-list/export/xlsx",
            json={"records": [record]},
        )
        assert invalid_archive.status_code == 409
        session.headers.pop("X-AYEC-Fiscal-Archive")

        rejected = session.post(
            f"{base_url}/api/desktop/service-list/export/xlsx",
            json={"records": [{**record, "unexpected": "field"}]},
        )
        assert rejected.status_code == 400
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)
