import sqlite3
import threading
from http.server import ThreadingHTTPServer

import requests
from pypdf import PdfReader

from tests.test_vendor_management_center import _setup_tenant
import Web_Arayuzu.Main as web_main


def test_stock_movement_history_delete_requires_password_and_only_soft_deletes(monkeypatch, tmp_path):
    _actor, tenant_path = _setup_tenant(monkeypatch, tmp_path)
    with sqlite3.connect(tenant_path) as conn:
        conn.execute("INSERT INTO stock_movements(id,description,is_deleted) VALUES(71,'Test movement',0)")
    server = ThreadingHTTPServer(("127.0.0.1", 0), web_main.AYECRequestHandler)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base_url = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        session = requests.Session()
        session.headers["Origin"] = base_url
        login = session.post(
            f"{base_url}/api/auth/login",
            json={"identifier": web_main.VENDOR_EMAIL, "password": "operator-test-pass", "tenant_id": "tenant-1"},
        )
        assert login.status_code == 200
        wrong = session.post(
            f"{base_url}/api/desktop/stock/movements/delete",
            json={"id": 71, "password": "incorrect"},
        )
        assert wrong.status_code == 403
        with sqlite3.connect(tenant_path) as conn:
            assert conn.execute("SELECT is_deleted FROM stock_movements WHERE id=71").fetchone()[0] == 0
        response = session.post(
            f"{base_url}/api/desktop/stock/movements/delete",
            json={"id": 71, "password": "operator-test-pass"},
        )
        assert response.status_code == 200
        assert response.json() == {"ok": True, "id": 71}
        with sqlite3.connect(tenant_path) as conn:
            assert conn.execute("SELECT is_deleted,deleted_at FROM stock_movements WHERE id=71").fetchone()[0] == 1
            assert conn.execute("SELECT COUNT(*) FROM stock_movements WHERE id=71").fetchone()[0] == 1
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)


def test_stock_movement_history_is_tenant_scoped_paginated_and_exported_as_xlsx(monkeypatch, tmp_path):
    _actor, tenant_path = _setup_tenant(monkeypatch, tmp_path)
    with sqlite3.connect(tenant_path) as conn:
        part_columns = {row[1] for row in conn.execute("PRAGMA table_info(parts)")}
        part_values = {"id": 901, "name": "History fixture", "part_name": "History fixture", "purchase_price": 12.5, "currency": "USD"}
        selected = {key: value for key, value in part_values.items() if key in part_columns}
        conn.execute(
            "INSERT INTO parts (" + ",".join('"' + key + '"' for key in selected) + ") VALUES (" + ",".join("?" for _ in selected) + ")",
            tuple(selected.values()),
        )
        movement_columns = {row[1] for row in conn.execute("PRAGMA table_info(stock_movements)")}
        for movement_id in range(1, 102):
            values = {
                "id": movement_id,
                "part_id": 901,
                "movement_type": "Giri\u015f" if movement_id % 2 == 0 else "\u00c7\u0131k\u0131\u015f",
                "amount": movement_id,
                "new_stock": movement_id,
                "description": "Movement " + str(movement_id),
                "created_at": "2026-01-01 " + str((movement_id - 1) // 60).zfill(2) + ":" + str((movement_id - 1) % 60).zfill(2) + ":00",
                "is_deleted": 0,
            }
            selected = {key: value for key, value in values.items() if key in movement_columns}
            conn.execute(
                "INSERT INTO stock_movements (" + ",".join('"' + key + '"' for key in selected) + ") VALUES (" + ",".join("?" for _ in selected) + ")",
                tuple(selected.values()),
            )

    server = ThreadingHTTPServer(("127.0.0.1", 0), web_main.AYECRequestHandler)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base_url = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        anonymous = requests.Session()
        anonymous.headers["Origin"] = base_url
        assert anonymous.get(f"{base_url}/api/desktop/stock/movements/export/pdf").status_code == 401
        session = requests.Session()
        session.headers["Origin"] = base_url
        login = session.post(
            f"{base_url}/api/auth/login",
            json={"identifier": web_main.VENDOR_EMAIL, "password": "operator-test-pass", "tenant_id": "tenant-1"},
        )
        assert login.status_code == 200
        page_two = session.get(f"{base_url}/api/desktop/stock/movements/history?page=2")
        assert page_two.status_code == 200
        assert page_two.json()["total"] == 101
        assert [row["id"] for row in page_two.json()["rows"]] == [1]
        assert page_two.json()["rows"][0]["part_name"] == "History fixture"
        filtered = session.get(
            f"{base_url}/api/desktop/stock/movements/history",
            params={"search": "Movement 10", "type": "Giri\u015f"},
        )
        assert filtered.status_code == 200
        assert all("10" in row["description"] and row["movement_type"] == "Giri\u015f" for row in filtered.json()["rows"])
        workbook_response = session.get(
            f"{base_url}/api/desktop/stock/movements/export/xlsx?page=2"
        )
        assert workbook_response.status_code == 200
        assert "stok_hareketleri_" in workbook_response.headers["Content-Disposition"]
        from io import BytesIO
        from openpyxl import load_workbook
        sheet = load_workbook(BytesIO(workbook_response.content), read_only=True).active
        assert list(next(sheet.iter_rows(values_only=True))) == [
            "ID", "Par\u00e7a", "\u0130\u015flem", "Miktar", "PB", "Birim Maliyet", "A\u00e7\u0131klama", "Tarih"
        ]
        exported_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        assert len(exported_rows) == 101
        assert {row[0] for row in exported_rows} == {str(movement_id) for movement_id in range(1, 102)}
        assert next(row for row in exported_rows if row[0] == "1") == (
            "1", "History fixture", "\u00c7\u0131k\u0131\u015f", "1.0", "USD", "12.50 $", "Movement 1", "01.01.2026 00:00"
        )
        pdf_response = session.get(
            f"{base_url}/api/desktop/stock/movements/export/pdf",
            params={"search": "Movement 10", "type": "Giri\u015f"},
        )
        assert pdf_response.status_code == 200
        assert pdf_response.headers["Content-Type"] == "application/pdf"
        assert pdf_response.content.startswith(b"%PDF-")
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf_response.content)).pages)
        assert "AYEC Pro - Stok Hareketleri" in pdf_text
        assert "Movement 10" in pdf_text
        assert "Movement 100" in pdf_text
        assert "Movement 101" not in pdf_text
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)
