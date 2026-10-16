from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.database import Database
from Web_Arayuzu import Main as web_main


ROOT = Path(__file__).resolve().parents[1]


def _project_with_units(path):
    database = Database(str(path))
    project_id = database.add_project({"name": "Unit Export Project"})
    root_id = database.add_project_unit({
        "project_id": project_id,
        "block_name": "A Block",
        "unit_no": "A-1",
        "status": "Devam Ediyor",
    })
    child_id = database.add_project_unit({
        "project_id": project_id,
        "block_name": "A Block",
        "unit_no": "A-1-Living",
        "parent_unit_id": root_id,
        "status": "Bekliyor",
        "unit_kind": "Oda",
    })
    assert database.add_unit_product({
        "project_id": project_id,
        "unit_id": root_id,
        "part_name": "Gateway",
        "part_code": "GW-1",
        "quantity": 1,
        "notes": "Root installation",
    })
    assert database.add_unit_product({
        "project_id": project_id,
        "unit_id": child_id,
        "part_name": "Sensor",
        "part_code": "SNS-2",
        "quantity": 2,
        "notes": "Room installation",
    })
    database.close()
    return project_id, root_id


def test_project_unit_xlsx_matches_desktop_two_sheet_report(monkeypatch, tmp_path):
    database_path = tmp_path / "unit-export.db"
    project_id, root_id = _project_with_units(database_path)
    monkeypatch.setattr(web_main, "active_db_path", lambda: database_path)

    payload = web_main.desktop_project_units_export_xlsx(project_id)
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Oda \u00d6zeti", "\u00dcr\u00fcn Detaylar\u0131"]
        summary = workbook["Oda \u00d6zeti"]
        assert list(summary.values)[0] == ("Blok", "Daire / Oda", "Durum", "Kurulu \u00dcr\u00fcn Say\u0131s\u0131")
        assert list(summary.values)[1] == ("A Block", "A-1", "Devam Ediyor", 2)
        details = list(workbook["\u00dcr\u00fcn Detaylar\u0131"].values)
        assert details[0] == ("Blok", "Daire / Oda", "\u00dcr\u00fcn Ad\u0131", "Kod", "Adet", "Teknik Not", "Eklenme Tarihi")
        assert {(row[2], row[3], row[4], row[5]) for row in details[1:]} == {
            ("Gateway", "GW-1", 1, "Root installation"),
            ("Sensor", "SNS-2", 2, "Room installation"),
        }
    finally:
        workbook.close()


def test_project_unit_csv_matches_desktop_product_detail_export(monkeypatch, tmp_path):
    database_path = tmp_path / "unit-export-csv.db"
    _project_with_units(database_path)
    monkeypatch.setattr(web_main, "active_db_path", lambda: database_path)

    payload = web_main.desktop_project_units_export_csv(1)
    assert payload.startswith(b"\xef\xbb\xbf")
    text = payload.decode("utf-8-sig")
    assert "Blok;Daire / Oda;\u00dcr\u00fcn Ad\u0131;Kod;Adet;Teknik Not;Eklenme Tarihi" in text
    assert "A Block;A-1;Gateway;GW-1;1;Root installation" in text
    assert "A Block;A-1-Living;Sensor;SNS-2;2;Room installation" in text


def test_project_unit_export_rejects_unknown_project(monkeypatch, tmp_path):
    database_path = tmp_path / "unit-export-missing.db"
    Database(str(database_path)).close()
    monkeypatch.setattr(web_main, "active_db_path", lambda: database_path)

    with pytest.raises(LookupError, match="Project was not found"):
        web_main.desktop_project_units_export_xlsx(999)


def test_project_unit_export_dialog_offers_xlsx_and_csv():
    app = (ROOT / "Web_Arayuzu" / "web" / "app.js").read_text(encoding="utf-8")
    start = app.index("if (unitExport)")
    end = app.index("if (unitBulk || unitEdit)", start)
    flow = app[start:end]

    assert 'value="xlsx"' in flow
    assert 'value="csv"' in flow
    assert "/units/export.${format}" in flow
    assert "credentials: \"same-origin\"" in flow
