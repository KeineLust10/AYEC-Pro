import io
import csv
import sqlite3
import threading
from datetime import datetime
from http.server import ThreadingHTTPServer

import requests
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.test_vendor_management_center import _setup_tenant
import Web_Arayuzu.Main as web_main


def test_accounting_ledger_xlsx_export_is_authenticated_and_formula_safe(monkeypatch, tmp_path):
    actor, tenant_path = _setup_tenant(monkeypatch, tmp_path)
    current_date = datetime.now().strftime("%Y-%m-%d")
    with sqlite3.connect(tenant_path) as conn:
        income_id = conn.execute(
            "INSERT INTO accounting(type,category,amount,description,date,customer_name) VALUES(?,?,?,?,?,?)",
            ("Gelir", "Servis", 120.5, "Needle income", current_date, "Test Customer"),
        ).lastrowid
        conn.execute(
            "INSERT INTO accounting(type,category,amount,description,date,customer_name) VALUES(?,?,?,?,?,?)",
            ("Gider", "Malzeme", 7.25, "Needle expense", current_date, "Test Vendor"),
        )
        conn.execute(
            "INSERT INTO accounting(type,category,amount,description,date) VALUES(?,?,?,?,?)",
            ("Gelir", "Eski", 88, "Needle old", "2020-01-01"),
        )
        conn.execute(
            "INSERT INTO accounting(type,category,amount,description,date) VALUES(?,?,?,?,?)",
            ("Gelir", "", 1, "=CMD()", current_date),
        )
        currency_id = conn.execute(
            "INSERT INTO currency_transactions(customer_id,transaction_type,amount,currency,exchange_rate,try_equivalent,description,tracking_no,created_at,current_balance) "
            "VALUES(NULL,?,?,?,?,?,?,?,?,?)",
            ("CREDIT", 20, "USD", 35, 700, "Needle currency", "FX-1", current_date + "T12:00:00", 0),
        ).lastrowid
        conn.commit()
    server = ThreadingHTTPServer(("127.0.0.1", 0), web_main.AYECRequestHandler)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base_url = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        anonymous = requests.Session()
        anonymous.headers["Origin"] = base_url
        endpoint = f"{base_url}/api/desktop/accounting/ledger/export/xlsx"
        assert anonymous.post(endpoint, json={"rows": []}).status_code == 401
        assert anonymous.post(f"{base_url}/api/desktop/accounting/ledger/export/pdf", json={"rows": []}).status_code == 401
        page_endpoint = f"{base_url}/api/desktop/accounting/ledger/page"
        assert anonymous.post(page_endpoint, json={"page": 0}).status_code == 401

        session = requests.Session()
        session.headers["Origin"] = base_url
        login = session.post(
            f"{base_url}/api/auth/login",
            json={"identifier": web_main.VENDOR_EMAIL, "password": "operator-test-pass", "tenant_id": "tenant-1"},
        )
        assert login.status_code == 200
        page_response = session.post(page_endpoint, json={"period": "year", "page": 0})
        assert page_response.status_code == 200
        page_payload = page_response.json()
        assert page_payload["page"] == 0 and page_payload["limit"] == 100
        assert isinstance(page_payload["rows"], list)
        assert session.post(page_endpoint, json={"page": -1}).status_code == 400
        response = session.post(
            endpoint,
            json={"period": "year", "filter": "INCOME", "search": "Needle", "rows": [{"Kayit ID": 17}]},
        )
        assert response.status_code == 200
        assert response.headers["Content-Type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        sheet = workbook.active
        assert sheet.max_row == 3
        exported = {
            (sheet.cell(row=index, column=1).value, sheet.cell(row=index, column=2).value): index
            for index in range(2, sheet.max_row + 1)
        }
        income_row = exported[(income_id, "accounting")]
        currency_row = exported[(currency_id, "currency")]
        assert sheet.cell(income_row, 4).value == "Gelir"
        assert sheet.cell(income_row, 5).value == "Servis - Needle income"
        assert sheet.cell(income_row, 10).value == 120.5
        assert sheet.cell(currency_row, 2).value == "currency"
        assert sheet.cell(currency_row, 4).value == "Tahsilat"
        assert sheet.cell(currency_row, 10).value == 700
        pdf_response = session.post(
            f"{base_url}/api/desktop/accounting/ledger/export/pdf",
            json={"period": "year", "filter": "INCOME", "search": "Needle"},
        )
        assert pdf_response.status_code == 200
        assert pdf_response.headers["Content-Type"] == "application/pdf"
        assert pdf_response.content.startswith(b"%PDF-")
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf_response.content)).pages)
        assert "AYEC Pro - Cari Hareketler" in pdf_text
        assert "Needle" in pdf_text and "income" in pdf_text
        assert "currency" in pdf_text
        assert "expense" not in pdf_text
        csv_endpoint = f"{base_url}/api/desktop/income-expense/export/csv"
        csv_response = session.post(
            csv_endpoint,
            json={"month": current_date[:7], "search": "Needle", "tab": "income"},
        )
        assert csv_response.status_code == 200
        csv_rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
        assert csv_rows[0] == ["Date", "Category", "Description", "Reference", "Amount TRY"]
        assert len(csv_rows) == 3
        csv_by_description = {row[2]: row for row in csv_rows[1:]}
        assert csv_by_description["Servis - Needle income"][1] == "Servis"
        assert csv_by_description["Needle currency"][1] == "Tahsilat"
        assert csv_by_description["Needle currency"][3] == "FX-1"
        assert csv_by_description["Needle currency"][4] == "700.0"
        formula_csv = session.post(
            csv_endpoint,
            json={"month": current_date[:7], "search": "=CMD()", "tab": "all"},
        )
        assert formula_csv.status_code == 200
        formula_rows = list(csv.reader(io.StringIO(formula_csv.content.decode("utf-8-sig"))))
        assert formula_rows[1][1] == "Gelir"
        assert formula_rows[1][3] == ""
        assert formula_rows[1][2] == "'=CMD()"
        expenses_csv = session.post(
            csv_endpoint,
            json={"month": current_date[:7], "search": "Needle", "tab": "expense"},
        )
        assert expenses_csv.status_code == 200
        assert len(list(csv.reader(io.StringIO(expenses_csv.content.decode("utf-8-sig"))))) == 2
        assert session.post(csv_endpoint, json={"month": "bad", "tab": "all"}).status_code == 400
        expenses = session.post(
            endpoint,
            json={"period": "year", "filter": "EXPENSE", "search": "Needle"},
        )
        assert expenses.status_code == 200
        expense_sheet = load_workbook(io.BytesIO(expenses.content), data_only=False).active
        assert expense_sheet.max_row == 2
        assert expense_sheet["D2"].value == "Gider"
        assert session.post(endpoint, json={"period": "invalid"}).status_code == 400
        for period in ("today", "last7", "week", "month", "fiscal", "year"):
            period_export = session.post(
                endpoint,
                json={"period": period, "filter": "ALL", "search": "Needle"},
            )
            assert period_export.status_code == 200
            period_sheet = load_workbook(io.BytesIO(period_export.content), data_only=False).active
            period_rows = {
                (period_sheet.cell(row=index, column=1).value, period_sheet.cell(row=index, column=2).value)
                for index in range(2, period_sheet.max_row + 1)
            }
            assert period_rows == {
                (income_id, "accounting"),
                (int(expense_sheet["A2"].value), "accounting"),
                (currency_id, "currency"),
            }

        safe = load_workbook(
            io.BytesIO(web_main.desktop_accounting_ledger_workbook([{"Aciklama": "=1+1"}])),
            data_only=False,
        ).active
        assert safe["E2"].value == "'=1+1"
        try:
            web_main.desktop_accounting_ledger_workbook([{}] * 100001)
        except ValueError:
            pass
        else:
            raise AssertionError("The desktop ledger export row limit must be enforced")

        snapshot_path = tmp_path / "accounting-archive.db"
        with sqlite3.connect(tenant_path) as source, sqlite3.connect(snapshot_path) as target:
            source.backup(target)
        archive = sqlite3.connect(snapshot_path)
        try:
            archive.execute(
                "INSERT INTO accounting(type,category,amount,description,date) VALUES(?,?,?,?,?)",
                ("Gelir", "Arsiv", 44, "Archived Only", current_date),
            )
            archive.commit()
        finally:
            archive.close()
        stored = web_main._store_support_backup_for_tenant(
            "tenant-1",
            int(actor["id"]),
            snapshot_path.read_bytes(),
            "archive_2026.db",
            "fiscal-rollover",
            "teknik_servis",
            "web-runtime",
            "web-app",
            deduplicate=False,
        )
        web_main.set_tenant_context(web_main._tenant_record_required("tenant-1"))
        archive_file, _, _ = web_main._fiscal_archive_file(actor, int(stored["backup_id"]))
        with sqlite3.connect(archive_file) as archive:
            assert archive.execute(
                "SELECT COUNT(*) FROM accounting WHERE description='Archived Only'"
            ).fetchone()[0] == 1
        web_main.clear_tenant_context()
        before = snapshot_path.read_bytes()
        archive_rows = session.get(
            f"{base_url}/api/desktop/table/accounting?limit=100",
            headers={"X-AYEC-Fiscal-Archive": str(stored["backup_id"])},
        )
        assert archive_rows.status_code == 200
        assert any(row.get("description") == "Archived Only" for row in archive_rows.json()["rows"]), archive_rows.json()
        archived = session.post(
            endpoint,
            headers={"X-AYEC-Fiscal-Archive": str(stored["backup_id"])},
            json={"period": "year", "filter": "ALL", "search": "Archived Only"},
        )
        assert archived.status_code == 200, archived.text
        archive_sheet = load_workbook(io.BytesIO(archived.content), data_only=False).active
        assert archive_sheet.max_row == 2
        assert archive_sheet["E2"].value == "Arsiv - Archived Only"
        archived_csv = session.post(
            csv_endpoint,
            headers={"X-AYEC-Fiscal-Archive": str(stored["backup_id"])},
            json={"month": current_date[:7], "search": "Archived Only", "tab": "all"},
        )
        assert archived_csv.status_code == 200
        assert len(list(csv.reader(io.StringIO(archived_csv.content.decode("utf-8-sig"))))) == 2
        assert snapshot_path.read_bytes() == before
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=2)


def test_accounting_ledger_page_uses_bounded_offset_pages(monkeypatch, tmp_path):
    actor, tenant_path = _setup_tenant(monkeypatch, tmp_path)
    current_date = datetime.now().strftime("%Y-%m-%d")
    with sqlite3.connect(tenant_path) as conn:
        conn.executemany(
            "INSERT INTO accounting(type,category,amount,description,date) VALUES(?,?,?,?,?)",
            [("Gelir", "Page", index + 1, f"PageAudit-{index:03d}", current_date) for index in range(101)],
        )
        conn.commit()
    web_main.set_tenant_context(web_main.tenant_by_id("tenant-1"))
    try:
        first = web_main.desktop_accounting_ledger_page(actor, {"period": "year", "search": "PageAudit", "page": 0})
        second = web_main.desktop_accounting_ledger_page(actor, {"period": "year", "search": "PageAudit", "page": 1})
        for invalid in (-1, True, 1000001, "1"):
            try:
                web_main.desktop_accounting_ledger_page(actor, {"page": invalid})
            except ValueError:
                pass
            else:
                raise AssertionError("Invalid ledger page values must be rejected")
    finally:
        web_main.clear_tenant_context()
    assert first["limit"] == 100
    assert first["page"] == 0 and len(first["rows"]) == 100
    assert first["has_more"] is True
    assert second["page"] == 1 and second["has_more"] is False and len(second["rows"]) == 1
    assert {row["id"] for row in first["rows"]}.isdisjoint({row["id"] for row in second["rows"]})
    assert second["rows"][0]["description"].startswith("PageAudit-")
