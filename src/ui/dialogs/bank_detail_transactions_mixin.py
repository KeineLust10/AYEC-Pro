# -*- coding: utf-8 -*-

from PyQt6.QtCore import QDate, Qt
from PyQt6.QtWidgets import (
    QFileDialog,
    QHBoxLayout,
    QMessageBox,
    QPushButton,
    QTableWidgetItem,
    QWidget,
)
from datetime import datetime
from src.ui.dialogs.bank_detail_dialog_components import AlertEditDialog, BankCardDialog
from src.utils.currency_helper import CurrencyHelper
from src.utils.design_system import DesignTokens
from src.utils.toast_notification import show_success, show_error
from src.utils.message_helper import show_info, show_question, show_warning
from src.utils.logger import logger
from src.utils.theme_colors import qc, tc, theme_qss


class BankDetailTransactionsMixin:
    def _display_currency(self):
        return CurrencyHelper.get_code(self.db)

    def _format_display_money(self, amount_try):
        return CurrencyHelper.format_from_try(
            amount_try,
            db=self.db,
            currency_code=self._display_currency(),
            include_try_reference=False,
        )

    def _build_tx_query(self, count_only=False, limit=None, offset=None):
        if not self.accounting_cols:
            return ("SELECT 0", [])
        conditions = []
        params = []

        if "bank_account_id" in self.accounting_cols:
            conditions.append("(bank_account_id = ? OR related_account_id = ?)")
            params.extend([self.acc_id, self.acc_id])

        text = self.txt_search.text().strip()
        if text:
            conditions.append("(description LIKE ? OR CAST(id AS TEXT) LIKE ?)")
            params.extend([f"%{text}%", f"%{text}%"])

        tx_type = self.cmb_tx_type.currentText()
        if tx_type != "Tümü":
            if tx_type in ["Gelir", "Gider", "Transfer"]:
                conditions.append("type = ?")
                params.append(tx_type)
            elif tx_type in ["Havale", "EFT", "Otomatik Ödeme"]:
                if "payment_method" in self.accounting_cols:
                    conditions.append("payment_method LIKE ?")
                    params.append(f"%{tx_type}%")
                else:
                    conditions.append("description LIKE ?")
                    params.append(f"%{tx_type}%")
            else:
                conditions.append("(category LIKE ? OR description LIKE ?)")
                params.extend([f"%{tx_type}%", f"%{tx_type}%"])

        if self.chk_tx_date.isChecked():
            start = self.tx_date_from.date().toString("yyyy-MM-dd")
            end = self.tx_date_to.date().toString("yyyy-MM-dd")
            conditions.append("date >= ? AND date <= ?")
            params.extend([start, end])

        if self.spin_amount_min.value() > 0:
            conditions.append("ABS(amount) >= ?")
            params.append(float(self.spin_amount_min.value()))
        if self.spin_amount_max.value() > 0:
            conditions.append("ABS(amount) <= ?")
            params.append(float(self.spin_amount_max.value()))

        where = " WHERE {clause}".format(clause=" AND ".join(conditions)) if conditions else ""
        if count_only:
            return ("SELECT COUNT(*) FROM accounting{where}".format(where=where), params)

        query = "SELECT * FROM accounting{where} ORDER BY date DESC, id DESC".format(where=where)
        if limit:
            query += " LIMIT ? OFFSET ?"
            params.extend([limit, offset or 0])
        return query, params

    def _load_transactions(self):
        try:
            count_query, count_params = self._build_tx_query(count_only=True)
            self.db.cursor.execute(count_query, count_params)
            total = int(self.db.cursor.fetchone()[0] or 0)
        except Exception:
            total = 0
        self.txn_total = total

        pages = max(1, int((total + self.txn_page_size - 1) / self.txn_page_size))
        if self.txn_page > pages:
            self.txn_page = pages
        offset = (self.txn_page - 1) * self.txn_page_size

        rows = []
        try:
            query, params = self._build_tx_query(limit=self.txn_page_size, offset=offset)
            self.db.cursor.execute(query, params)
            rows = self.db.cursor.fetchall()
        except Exception:
            rows = []

        self.table_tx.setRowCount(len(rows))
        for i, row in enumerate(rows):
            data = {self.accounting_cols[idx]: row[idx] for idx in range(min(len(self.accounting_cols), len(row)))} if self.accounting_cols else {}
            date_val = data.get("date") or "—"
            created = data.get("created_at") or ""
            time_val = ""
            if created and isinstance(created, str) and " " in created:
                time_val = created.split(" ")[1][:5]
            ref_val = data.get("id")
            tx_type = data.get("type") or ""
            desc = data.get("description") or ""
            amount = float(data.get("amount") or 0)
            pay_method = data.get("payment_method") or "—"
            related = data.get("related_account_id")
            customer_name = data.get("customer_name")
            counterparty = "—"
            if related:
                counterparty = self.bank_accounts.get(int(related), f"#{related}")
            elif customer_name:
                counterparty = str(customer_name)

            self.table_tx.setItem(i, 0, QTableWidgetItem(str(date_val)))
            self.table_tx.setItem(i, 1, QTableWidgetItem(str(time_val)))
            self.table_tx.setItem(i, 2, QTableWidgetItem(str(ref_val)))
            self.table_tx.setItem(i, 3, QTableWidgetItem(str(tx_type)))
            self.table_tx.setItem(i, 4, QTableWidgetItem(str(desc)))

            amt_item = QTableWidgetItem(self._format_display_money(amount))
            if tx_type == "Gelir" or (tx_type == "Transfer" and str(data.get("category")).find("Giriş") >= 0):
                amt_item.setForeground(qc("success"))
            elif tx_type == "Gider" or (tx_type == "Transfer" and str(data.get("category")).find("Çıkış") >= 0):
                amt_item.setForeground(qc("danger"))
            self.table_tx.setItem(i, 5, amt_item)

            self.table_tx.setItem(i, 6, QTableWidgetItem(str(pay_method)))
            self.table_tx.setItem(i, 7, QTableWidgetItem(str(counterparty)))

            self.table_tx.item(i, 0).setData(Qt.ItemDataRole.UserRole, data)

        self.lbl_tx_count.setText(f"{total} kayıt")
        self.lbl_page.setText(f"{self.txn_page} / {pages}")
        self.btn_prev.setEnabled(self.txn_page > 1)
        self.btn_next.setEnabled(self.txn_page < pages)

    def _change_page(self, delta):
        self.txn_page = max(1, self.txn_page + delta)
        self._load_transactions()

    def _change_page_size(self):
        try:
            self.txn_page_size = int(self.cmb_page_size.currentText())
        except Exception:
            self.txn_page_size = 25
        self.txn_page = 1
        self._load_transactions()

    def _apply_filters(self):
        self.txn_page = 1
        self._load_transactions()

    def _reset_filters(self):
        self.txt_search.clear()
        self.cmb_tx_type.setCurrentText("Tümü")
        self.chk_tx_date.setChecked(False)
        self.tx_date_from.setDate(QDate.currentDate().addMonths(-1))
        self.tx_date_to.setDate(QDate.currentDate())
        self.spin_amount_min.setValue(0)
        self.spin_amount_max.setValue(0)
        self.txn_page = 1
        self._load_transactions()

    def _open_tx_detail(self, item):
        row = item.row()
        data = self.table_tx.item(row, 0).data(Qt.ItemDataRole.UserRole)
        if not data:
            return
        try:
            from src.ui.pages.accounting_page import TransactionDetailsDialog
            entry = {
                "source": "accounting",
                "id": data.get("id"),
                "type_label": f"{data.get('type') or ''} / {data.get('category') or ''}".strip(" /"),
                "date": data.get("date"),
                "description": data.get("description"),
                "currency": data.get("currency") or "TRY",
                "amount": float(data.get("amount") or 0.0),
                "rate": float(data.get("exchange_rate") or 1.0),
                "amount_try": float(data.get("try_equivalent") or data.get("amount") or 0.0),
                "status": "—",
                "customer_id": data.get("customer_id"),
                "customer_name": data.get("customer_name"),
                "tracking_no": data.get("tracking_no"),
                "balance": None,
                "bank_account_id": data.get("bank_account_id"),
                "ref_no": data.get("ref_no"),
                "selected_services": data.get("selected_services"),
            }
            TransactionDetailsDialog(self.db, entry, self).exec()
        except Exception as e:
            logger.debug(f"Bank transaction detail dialog open failed: {e}")
            return

    def _export_transactions(self, kind):
        rows = self._fetch_all_transactions()
        if not rows:
            show_warning(self, "Uyarı", "Dışa aktarılacak kayıt bulunamadı.")
            return
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        if kind == "excel":
            path, _ = QFileDialog.getSaveFileName(self, "Excel'e Aktar", f"banka_hareketleri_{timestamp}.xlsx", "Excel Files (*.xlsx)")
            if not path:
                return
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment
            except ImportError:
                show_warning(self, "Uyarı", "Excel aktarımı için openpyxl kurulu olmalı.")
                return
            wb = Workbook()
            ws = wb.active
            ws.title = "Hareketler"
            headers = ["Tarih", "Saat", "Ref", "Tür", "Açıklama", "Tutar", "Yöntem", "Gönderen/Alan"]
            ws.append(headers)
            for row in rows:
                ws.append(row)
            header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
            for cell in ws[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            wb.save(path)
            show_info(self, "Başarılı", "Excel dosyası kaydedildi.")
            return

        path, _ = QFileDialog.getSaveFileName(self, "PDF'e Aktar", f"banka_hareketleri_{timestamp}.pdf", "PDF Files (*.pdf)")
        if not path:
            return
        try:
            from src.ui.utils.background_task import run_cancellable_task

            text_color = tc("text")
            selection_text = tc("selection_text")
            border_color = tc("border")

            def produce_pdf(is_cancelled):
                from reportlab.lib.pagesizes import A4, landscape
                from reportlab.platypus import (
                    SimpleDocTemplate,
                    Table,
                    TableStyle,
                    Paragraph,
                    Spacer,
                )
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.lib import colors

                if is_cancelled():
                    return None
                doc = SimpleDocTemplate(
                    path,
                    pagesize=landscape(A4),
                    rightMargin=24,
                    leftMargin=24,
                    topMargin=24,
                    bottomMargin=24,
                )
                styles = getSampleStyleSheet()
                story = [
                    Paragraph("Banka Hesap Hareketleri", styles["Heading1"]),
                    Spacer(1, 12),
                ]
                data = [
                    [
                        "Tarih",
                        "Saat",
                        "Ref",
                        "T\u00fcr",
                        "A\u00e7\u0131klama",
                        "Tutar",
                        "Y\u00f6ntem",
                        "G\u00f6nderen/Alan",
                    ]
                ]
                data.extend(rows)
                table = Table(data, repeatRows=1)
                table.setStyle(
                    TableStyle(
                        [
                            (
                                "BACKGROUND",
                                (0, 0),
                                (-1, 0),
                                colors.HexColor(text_color),
                            ),
                            (
                                "TEXTCOLOR",
                                (0, 0),
                                (-1, 0),
                                colors.HexColor(selection_text),
                            ),
                            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                            (
                                "GRID",
                                (0, 0),
                                (-1, -1),
                                0.25,
                                colors.HexColor(border_color),
                            ),
                            ("FONTSIZE", (0, 0), (-1, -1), 8),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ]
                    )
                )
                story.append(table)
                doc.build(story)
                return path

            run_cancellable_task(
                owner=self,
                title="Banka hareketleri PDF olusturuluyor",
                target=produce_pdf,
                on_success=lambda result_path: (
                    show_info(
                        self,
                        "Ba\u015far\u0131l\u0131",
                        "PDF dosyas\u0131 kaydedildi.",
                    )
                    if result_path
                    else None
                ),
                on_error=lambda message: show_warning(
                    self,
                    "Hata",
                    f"PDF olu\u015fturulamad\u0131: {message}",
                ),
                output_path=path,
            )
        except Exception as e:
            show_warning(
                self,
                "Hata",
                f"PDF olu\u015fturulamad\u0131: {e}",
            )

    def _fetch_all_transactions(self):
        try:
            query, params = self._build_tx_query(limit=None)
            self.db.cursor.execute(query, params)
            rows = self.db.cursor.fetchall()
        except Exception as e:
            logger.warning(f"Bank transactions fetch-all failed: {e}")
            rows = []
        output = []
        for row in rows:
            data = {self.accounting_cols[idx]: row[idx] for idx in range(min(len(self.accounting_cols), len(row)))} if self.accounting_cols else {}
            date_val = data.get("date") or "—"
            created = data.get("created_at") or ""
            time_val = ""
            if created and isinstance(created, str) and " " in created:
                time_val = created.split(" ")[1][:5]
            ref_val = data.get("id")
            tx_type = data.get("type") or ""
            desc = data.get("description") or ""
            amount = float(data.get("amount") or 0)
            pay_method = data.get("payment_method") or "—"
            related = data.get("related_account_id")
            customer_name = data.get("customer_name")
            counterparty = "—"
            if related:
                counterparty = self.bank_accounts.get(int(related), f"#{related}")
            elif customer_name:
                counterparty = str(customer_name)
            output.append([
                str(date_val),
                str(time_val),
                str(ref_val),
                str(tx_type),
                str(desc),
                self._format_display_money(amount),
                str(pay_method),
                str(counterparty),
            ])
        return output

    def _compute_summary_totals(self):
        totals = {
            "income": 0.0,
            "expense": 0.0,
            "net": 0.0,
            "week_income": 0.0,
            "week_expense": 0.0,
            "month_income": 0.0,
            "month_expense": 0.0,
        }
        if not self.acc_id or "bank_account_id" not in self.accounting_cols:
            return totals
        try:
            self.db.cursor.execute(
                "SELECT type, category, amount, date FROM accounting WHERE bank_account_id=? OR related_account_id=?",
                (self.acc_id, self.acc_id),
            )
            rows = self.db.cursor.fetchall()
        except Exception as e:
            logger.warning(f"Bank transactions summary query failed: {e}")
            rows = []

        def is_recent(date_str, days):
            try:
                d = datetime.strptime(date_str, "%Y-%m-%d")
                return (datetime.now() - d).days <= days
            except Exception:
                return False

        for row in rows:
            tx_type, category, amount, date_str = row[0], row[1], float(row[2] or 0), row[3]
            if tx_type == "Gelir" or (tx_type == "Transfer" and str(category).find("Giriş") >= 0):
                totals["income"] += amount
                if is_recent(date_str, 7):
                    totals["week_income"] += amount
                if is_recent(date_str, 30):
                    totals["month_income"] += amount
            elif tx_type == "Gider" or (tx_type == "Transfer" and str(category).find("Çıkış") >= 0):
                totals["expense"] += abs(amount)
                if is_recent(date_str, 7):
                    totals["week_expense"] += abs(amount)
                if is_recent(date_str, 30):
                    totals["month_expense"] += abs(amount)
        totals["net"] = totals["income"] - totals["expense"]
        return totals

    def _get_category_distribution(self):
        if not self.acc_id or "bank_account_id" not in self.accounting_cols:
            return []
        try:
            self.db.cursor.execute(
                "SELECT category, SUM(amount) FROM accounting WHERE type='Gider' AND (bank_account_id=? OR related_account_id=?) GROUP BY category ORDER BY SUM(amount) DESC LIMIT 8",
                (self.acc_id, self.acc_id),
            )
            rows = self.db.cursor.fetchall()
            return [(str(r[0] or "Diğer"), float(r[1] or 0)) for r in rows]
        except Exception as e:
            logger.warning(f"Bank transactions category distribution failed: {e}")
            return []

    def _load_alerts(self):
        alerts = []
        try:
            self.db.cursor.execute("SELECT alert_type, trigger_time, is_enabled FROM scheduled_alerts ORDER BY trigger_time ASC")
            rows = self.db.cursor.fetchall()
            for alert_type, trigger_time, is_enabled in rows:
                label = str(alert_type or "").replace("_", " ").title()
                status = "Aktif" if int(is_enabled or 0) == 1 else "Pasif"
                alerts.append({"label": label, "time": trigger_time, "status": status, "alert_type": alert_type, "is_enabled": int(is_enabled or 0)})
        except Exception as e:
            logger.warning(f"Bank alerts load failed: {e}")
        if not alerts:
            alerts.append({"label": "Kayıt bulunamadı", "time": "—", "status": "—", "alert_type": None, "is_enabled": 0})
        return alerts

    def _load_cards(self):
        if not self.acc_id:
            return []
        try:
            self.db.create_finance_tables()
            rows = self.db.get_bank_cards(self.acc_id)
            cols = self.card_cols
            if cols:
                return [{cols[i]: row[i] for i in range(min(len(cols), len(row)))} for row in rows]
            return [dict(r) for r in rows]
        except Exception:
            return []

    def _refresh_alerts(self):
        alerts = self._load_alerts()
        self.alert_table.setRowCount(len(alerts))
        for i, row in enumerate(alerts):
            label_item = QTableWidgetItem(row.get("label", "—"))
            label_item.setData(Qt.ItemDataRole.UserRole, row)
            self.alert_table.setItem(i, 0, label_item)
            self.alert_table.setItem(i, 1, QTableWidgetItem(row.get("time", "—")))
            self.alert_table.setItem(i, 2, QTableWidgetItem(row.get("status", "—")))
            if row.get("alert_type"):
                btn_edit = QPushButton("Düzenle")
                btn_edit.setStyleSheet(theme_qss(DesignTokens.get_button_qss("outline")))
                btn_edit.setCursor(Qt.CursorShape.PointingHandCursor)
                btn_edit.clicked.connect(lambda _, r=row: self._open_alert_editor(r))
                cell = QWidget()
                layout = QHBoxLayout(cell)
                layout.setContentsMargins(0, 0, 0, 0)
                layout.addStretch()
                layout.addWidget(btn_edit)
                layout.addStretch()
                self.alert_table.setCellWidget(i, 3, cell)
            else:
                self.alert_table.setItem(i, 3, QTableWidgetItem("—"))

    def _refresh_cards(self):
        cards = self._load_cards()
        self.cards_table.clearSpans()
        self.cards_table.setRowCount(len(cards) if cards else 1)
        if not cards:
            self.cards_table.setItem(0, 0, QTableWidgetItem("Kayıt bulunamadı"))
            self.cards_table.setSpan(0, 0, 1, 5)
            return
        for i, card in enumerate(cards):
            name = card.get("card_name") or "Kart"
            last4 = str(card.get("card_last4") or "").strip()
            label = f"{name} ••••{last4}" if last4 else name
            limit_val = float(card.get("card_limit") or 0)
            debt_val = float(card.get("current_debt") or 0)
            status = "Aktif" if int(card.get("is_active") or 0) == 1 else "Pasif"
            item = QTableWidgetItem(label)
            item.setData(Qt.ItemDataRole.UserRole, card)
            self.cards_table.setItem(i, 0, item)
            self.cards_table.setItem(i, 1, QTableWidgetItem(self._format_display_money(limit_val)))
            self.cards_table.setItem(i, 2, QTableWidgetItem(self._format_display_money(debt_val)))
            self.cards_table.setItem(i, 3, QTableWidgetItem(status))
            btn_edit = QPushButton("Düzenle")
            btn_edit.setStyleSheet(theme_qss(DesignTokens.get_button_qss("outline")))
            btn_edit.setCursor(Qt.CursorShape.PointingHandCursor)
            btn_edit.clicked.connect(lambda _, c=card: self._open_edit_card_dialog(c))
            btn_del = QPushButton("Sil")
            btn_del.setStyleSheet(theme_qss(DesignTokens.get_button_qss("secondary")))
            btn_del.setCursor(Qt.CursorShape.PointingHandCursor)
            btn_del.clicked.connect(lambda _, c=card: self._delete_card(c))
            cell = QWidget()
            layout = QHBoxLayout(cell)
            layout.setContentsMargins(0, 0, 0, 0)
            layout.addWidget(btn_edit)
            layout.addWidget(btn_del)
            self.cards_table.setCellWidget(i, 4, cell)

    def _open_add_card_dialog(self):
        dialog = BankCardDialog(self.db, self.acc_id, parent=self)
        if dialog.exec():
            self._refresh_cards()

    def _open_edit_card_dialog(self, card):
        dialog = BankCardDialog(self.db, self.acc_id, card_data=card, parent=self)
        if dialog.exec():
            self._refresh_cards()

    def _delete_card(self, card):
        card_id = card.get("id")
        if not card_id:
            return
        res = show_question(self, "Onay", "Kart silinsin mi")
        if res != QMessageBox.StandardButton.Yes:
            return
        if self.db.delete_bank_card(card_id):
            self._refresh_cards()

    def _open_alert_from_table(self, item):
        row = item.row()
        data = self.alert_table.item(row, 0).data(Qt.ItemDataRole.UserRole)
        if data:
            self._open_alert_editor(data)

    def _open_alert_editor(self, alert):
        if not alert.get("alert_type"):
            return
        dialog = AlertEditDialog(self.db, alert, parent=self)
        if dialog.exec():
            self._refresh_alerts()

