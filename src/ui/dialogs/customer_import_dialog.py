# -*- coding: utf-8 -*-
import csv
import itertools
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
)

from src.utils.toast_notification import show_error, show_success, show_warning


class CustomerImportDialog(QDialog):
    """Preview and import customer rows from a spreadsheet or CSV file."""

    HEADER_SCAN_LIMIT = 20
    PREVIEW_ROW_LIMIT = 200
    IMPORT_BATCH_SIZE = 50

    COLUMNS = (
        ("name", "M\u00fc\u015fteri"),
        ("phone", "Telefon"),
        ("email", "E-posta"),
        ("company_name", "Firma"),
        ("address", "Adres"),
        ("tax_no", "Vergi No"),
        ("city", "Sehir"),
        ("type", "Tur"),
        ("notes", "Not"),
    )

    HEADER_ALIASES = {
        "name": {
            "musteri", "musteri adi", "ad soyad", "ad", "isim", "name",
            "unvan", "firma adi", "firma", "company",
        },
        "phone": {"telefon", "telefon no", "tel", "gsm", "cep", "phone"},
        "email": {"eposta", "e posta", "email", "mail"},
        "company_name": {"firma", "firma adi", "sirket", "kurum", "company"},
        "address": {"adres", "address"},
        "tax_no": {"vergi no", "vergi numarasi", "tax no", "tax number"},
        "city": {"sehir", "il", "city"},
        "type": {"tur", "musteri turu", "tip", "type"},
        "notes": {"not", "aciklama", "notes", "description"},
    }

    def __init__(self, db, source_path, parent=None):
        super().__init__(parent)
        self.db = db
        self.source_path = Path(source_path)
        self.rows = []
        self._mapping = {}
        self.setWindowTitle("M\u00fc\u015fteri I\u00e7e Aktar\u0131m \u00d6nizleme")
        self.resize(1120, 670)
        self._build_ui()
        self._load_rows()

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(18, 18, 18, 18)
        root.setSpacing(12)

        title = QLabel("M\u00fc\u015fteri I\u00e7e Aktar\u0131m \u00d6nizleme")
        title.setStyleSheet("font-size: 20px; font-weight: 700;")
        root.addWidget(title)

        self.info_label = QLabel()
        self.info_label.setWordWrap(True)
        root.addWidget(self.info_label)

        self.table = QTableWidget(0, len(self.COLUMNS))
        self.table.setHorizontalHeaderLabels([label for _, label in self.COLUMNS])
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.DoubleClicked)
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.verticalHeader().setVisible(False)
        root.addWidget(self.table, 1)

        actions = QHBoxLayout()
        actions.addStretch()
        cancel_button = QPushButton("Vazgec")
        cancel_button.clicked.connect(self.reject)
        import_button = QPushButton("Kaydet ve I\u00e7eri Aktar")
        import_button.setDefault(True)
        import_button.clicked.connect(self.save_rows)
        actions.addWidget(cancel_button)
        actions.addWidget(import_button)
        root.addLayout(actions)

    @staticmethod
    def _clean_header(value):
        text = str(value or "").strip().lower().replace("\u0131", "i")
        text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
        return re.sub(r"[^a-z0-9]+", " ", text).strip()

    def _iter_raw_rows(self):
        suffix = self.source_path.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(self.source_path, read_only=True, data_only=True)
            try:
                yield from workbook.active.iter_rows(values_only=True)
            finally:
                workbook.close()
            return
        if suffix == ".csv":
            last_error = None
            for encoding in ("utf-8-sig", "utf-8", "cp1254", "latin-1"):
                try:
                    with self.source_path.open("r", encoding=encoding, newline="") as handle:
                        yield from csv.reader(handle)
                    return
                except UnicodeDecodeError as exc:
                    last_error = exc
            raise last_error or ValueError("CSV could not be read")
        raise ValueError("Desteklenmeyen dosya turu")

    def _source_mapping_and_rows(self):
        raw_rows = self._iter_raw_rows()
        candidates = list(itertools.islice(raw_rows, self.HEADER_SCAN_LIMIT))
        if not candidates:
            raise ValueError("Dosyada aktarilacak satir yok.")

        def header_score(row):
            normalized = {self._clean_header(cell) for cell in row if cell is not None}
            return sum(
                bool(normalized.intersection(aliases))
                for aliases in self.HEADER_ALIASES.values()
            )

        header_index, _ = max(
            enumerate(candidates),
            key=lambda item: (header_score(item[1]), sum(bool(cell) for cell in item[1])),
        )
        return self._resolve_columns(candidates[header_index]), itertools.chain(
            candidates[header_index + 1:], raw_rows
        )

    def _resolve_columns(self, headers):
        normalized = [self._clean_header(value) for value in headers]
        mapping = {}
        for field, aliases in self.HEADER_ALIASES.items():
            for index, header in enumerate(normalized):
                if header in aliases:
                    mapping[field] = index
                    break
        if "name" not in mapping and normalized:
            mapping["name"] = 0
        return mapping

    @staticmethod
    def _cell(values, index):
        if index is None or index >= len(values) or values[index] is None:
            return ""
        return str(values[index]).strip()

    def _load_rows(self):
        try:
            mapping, source_rows = self._source_mapping_and_rows()
        except Exception as exc:
            show_error(self, f"Dosya okunamadi: {exc}")
            return

        self._mapping = mapping
        self.rows = []
        for source_index, values in enumerate(source_rows):
            row = {field: self._cell(values, mapping.get(field)) for field, _ in self.COLUMNS}
            if row["name"]:
                self.rows.append((source_index, row))
            if len(self.rows) >= self.PREVIEW_ROW_LIMIT:
                break

        if not self.rows:
            show_warning(self, "Dosyada aktarilacak musteri yok.")
            return

        matched = ", ".join(field for field, _ in self.COLUMNS if field in mapping)
        self.info_label.setText(
            f"Kaynak: {self.source_path.name} | Onizleme: ilk {len(self.rows)} satir | "
            f"Eslesen alanlar: {matched or 'ad otomatik secildi'}"
        )
        self.table.setRowCount(len(self.rows))
        for row_index, (source_index, row) in enumerate(self.rows):
            for column_index, (field, _) in enumerate(self.COLUMNS):
                item = QTableWidgetItem(row[field])
                if column_index == 0:
                    item.setData(Qt.ItemDataRole.UserRole, source_index)
                self.table.setItem(row_index, column_index, item)

    def _preview_overrides(self):
        rows = {}
        for row_index in range(self.table.rowCount()):
            row = {}
            for column_index, (field, _) in enumerate(self.COLUMNS):
                item = self.table.item(row_index, column_index)
                row[field] = item.text().strip() if item else ""
            if row["name"]:
                first_item = self.table.item(row_index, 0)
                source_index = first_item.data(Qt.ItemDataRole.UserRole) if first_item else None
                if source_index is not None:
                    rows[int(source_index)] = row
        return rows

    def _iter_import_rows(self, overrides):
        mapping, source_rows = self._source_mapping_and_rows()
        for source_index, values in enumerate(source_rows):
            row = overrides.get(source_index)
            if row is None:
                row = {field: self._cell(values, mapping.get(field)) for field, _ in self.COLUMNS}
            if row["name"]:
                yield row

    def save_rows(self):
        overrides = self._preview_overrides()
        if not overrides:
            show_warning(self, "Kaydedilecek musteri yok.")
            return

        saved_count = 0
        failures = []
        try:
            import_rows = self._iter_import_rows(overrides)
            for row_number, row in enumerate(import_rows, start=1):
                payload = {
                    "name": row["name"],
                    "phone": row["phone"],
                    "email": row["email"],
                    "company_name": row["company_name"],
                    "address": row["address"],
                    "tax_no": row["tax_no"],
                    "city": row["city"],
                    "type": row["type"] or "Bireysel",
                    "notes": row["notes"],
                    "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
                try:
                    customer_id = self.db.add_customer(payload)
                    if not customer_id:
                        detail = str(getattr(self.db, "_last_customer_error", "")).strip()
                        raise RuntimeError(detail or "add_customer returned no identifier")
                except Exception as exc:
                    failures.append(f"{row_number}: {exc}")
                    continue
                saved_count += 1
                if row_number % self.IMPORT_BATCH_SIZE == 0:
                    QApplication.processEvents()
        except Exception as exc:
            show_error(self, f"Dosya aktarimi baslatilamadi: {exc}")
            return

        if not saved_count:
            detail = failures[0] if failures else ""
            show_error(self, f"Musteriler kaydedilemedi. {detail}".strip())
            return

        message = f"Aktarma tamamlandi. Basarili: {saved_count} | Hatali: {len(failures)}"
        if failures:
            message += f"\nIlk hata: {failures[0]}"
            show_warning(self, message)
        else:
            show_success(self, message)
        self.accept()
