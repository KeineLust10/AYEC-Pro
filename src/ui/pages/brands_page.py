# -*- coding: utf-8 -*-

import html
from pathlib import Path

from PyQt6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QGridLayout,
    QLabel,
    QPushButton,
    QFrame,
    QTableWidget,
    QTableWidgetItem,
    QHeaderView,
    QLineEdit,
    QAbstractItemView,
    QFileDialog,
    QCheckBox,
    QStyledItemDelegate,
    QStyleOptionViewItem,
    QApplication,
    QStyle,
)
from PyQt6.QtCore import Qt, QSize
from PyQt6.QtGui import QColor, QFont, QPainter, QPen, QBrush, QTextDocument, QIcon, QPixmap, QPalette
from PyQt6.QtCore import QRect, QRectF
from PyQt6.QtPrintSupport import QPrintDialog, QPrinter

from src.utils.theme_colors import theme_qss
from src.utils.design_system import DesignTokens
from src.utils.toast_notification import show_success, show_error
from src.ui.widgets.modern_inputs import ModernComboBox
from src.ui.widgets.reference_status_delegate import ReferenceStatusDelegate
from src.utils.theme_colors import qc


def _toolbar_icon(kind):
    colors = {"catalog": "#38bdf8", "excel": "#22c55e", "pdf": "#fb7185",
              "print": "#cbd5e1", "add": "#ffffff"}
    color = QColor(colors[kind])
    pixmap = QPixmap(20, 20)
    pixmap.fill(Qt.GlobalColor.transparent)
    painter = QPainter(pixmap)
    painter.setRenderHint(QPainter.RenderHint.Antialiasing)
    painter.setPen(QPen(color, 1.8))
    painter.setBrush(QBrush(Qt.BrushStyle.NoBrush))
    if kind == "catalog":
        painter.drawArc(QRectF(3, 3, 14, 14), 30 * 16, 285 * 16)
        painter.drawLine(14, 2, 17, 4)
        painter.drawLine(14, 2, 13, 6)
    elif kind in {"excel", "pdf"}:
        painter.drawRoundedRect(QRectF(4, 2, 12, 16), 2, 2)
        painter.drawLine(7, 7, 13, 7)
        painter.drawLine(7, 10, 13, 10)
        painter.drawLine(7, 13, 13, 13)
    elif kind == "print":
        painter.drawRect(QRectF(5, 2, 10, 5))
        painter.drawRoundedRect(QRectF(3, 7, 14, 8), 2, 2)
        painter.drawRect(QRectF(5, 12, 10, 6))
        painter.drawLine(7, 15, 13, 15)
    else:
        painter.drawEllipse(QRectF(2, 2, 16, 16))
        painter.drawLine(10, 5, 10, 15)
        painter.drawLine(5, 10, 15, 10)
    painter.end()
    return QIcon(pixmap)


class _ServiceCountItem(QTableWidgetItem):
    def __init__(self, count, unit="Fi\u015f"):
        super().__init__('-' if count is None else '{} {}'.format(count, unit))
        self.setData(Qt.ItemDataRole.UserRole, -1 if count is None else count)

    def __lt__(self, other):
        return self.data(Qt.ItemDataRole.UserRole) < other.data(Qt.ItemDataRole.UserRole)


class _ReferenceSortItem(QTableWidgetItem):
    def __lt__(self, other):
        left = self.data(Qt.ItemDataRole.UserRole)
        right = other.data(Qt.ItemDataRole.UserRole)
        return str(left or self.text()).casefold() < str(right or other.text()).casefold()


class _DeviceGroupDelegate(QStyledItemDelegate):
    @staticmethod
    def _label_color(option):
        return option.palette.color(QPalette.ColorRole.Text)

    def paint(self, painter, option, index):
        base = QStyleOptionViewItem(option)
        self.initStyleOption(base, index)
        base.text = ""
        super().paint(painter, base, index)
        value = str(index.data(Qt.ItemDataRole.UserRole) or "")
        glyph, color = BrandsPage._group_glyph(value)
        rect = option.rect
        badge = QRect(rect.left() + 24, rect.top() + (rect.height() - 34) // 2, 34, 34)
        painter.save()
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setPen(QPen(QColor(color), 1))
        painter.setBrush(QBrush(QColor(color).darker(420)))
        painter.drawRoundedRect(badge, 6, 6)
        painter.setPen(QColor(color))
        painter.setBrush(QBrush(Qt.BrushStyle.NoBrush))
        painter.setPen(QPen(QColor(color), 1.7, Qt.PenStyle.SolidLine, Qt.PenCapStyle.RoundCap, Qt.PenJoinStyle.RoundJoin))
        cx = badge.center().x()
        cy = badge.center().y()
        if "ayd\u0131nlatma" in value.casefold():
            painter.drawEllipse(QRectF(cx - 6, cy - 9, 12, 12))
            painter.drawLine(cx - 4, cy + 4, cx + 4, cy + 4)
            painter.drawLine(cx - 3, cy + 7, cx + 3, cy + 7)
            painter.drawLine(cx - 9, cy - 3, cx - 11, cy - 3)
            painter.drawLine(cx + 9, cy - 3, cx + 11, cy - 3)
            painter.drawLine(cx, cy - 13, cx, cy - 15)
        elif "ev merkezi" in value.casefold() or "kamera" in value.casefold() or "priz" in value.casefold():
            painter.drawLine(cx - 9, cy - 1, cx, cy - 9)
            painter.drawLine(cx, cy - 9, cx + 9, cy - 1)
            painter.drawLine(cx - 6, cy - 2, cx - 6, cy + 8)
            painter.drawLine(cx + 6, cy - 2, cx + 6, cy + 8)
            painter.drawLine(cx - 6, cy + 8, cx + 6, cy + 8)
            painter.drawLine(cx - 1, cy + 8, cx - 1, cy + 2)
            painter.drawLine(cx + 2, cy + 8, cx + 2, cy + 2)
        elif "saat" in value.casefold():
            painter.drawLine(cx - 3, cy - 10, cx - 3, cy - 6)
            painter.drawLine(cx + 3, cy - 10, cx + 3, cy - 6)
            painter.drawLine(cx - 3, cy + 6, cx - 3, cy + 10)
            painter.drawLine(cx + 3, cy + 6, cx + 3, cy + 10)
            painter.drawRoundedRect(QRectF(cx - 7, cy - 7, 14, 14), 5, 5)
            painter.drawLine(cx, cy - 4, cx, cy)
            painter.drawLine(cx, cy, cx + 3, cy + 2)
        elif "telefon" in value.casefold():
            painter.drawRoundedRect(QRectF(cx - 6, cy - 10, 12, 20), 2, 2)
            painter.drawLine(cx - 2, cy - 7, cx + 2, cy - 7)
            painter.drawEllipse(QRectF(cx - 1, cy + 6, 2, 2))
        elif "laptop" in value.casefold() or "diz\u00fcst\u00fc" in value.casefold():
            painter.drawRoundedRect(QRectF(cx - 8, cy - 8, 16, 12), 1, 1)
            painter.drawLine(cx - 11, cy + 7, cx + 11, cy + 7)
            painter.drawLine(cx - 8, cy + 4, cx - 10, cy + 7)
            painter.drawLine(cx + 8, cy + 4, cx + 10, cy + 7)
        else:
            painter.drawRoundedRect(QRectF(cx - 9, cy - 8, 18, 13), 2, 2)
            painter.drawLine(cx, cy + 5, cx, cy + 9)
            painter.drawLine(cx - 5, cy + 9, cx + 5, cy + 9)
        painter.setPen(self._label_color(option))
        painter.setFont(QFont("Inter", 10, QFont.Weight.DemiBold))
        painter.drawText(QRect(badge.right() + 12, rect.top(), rect.right() - badge.right() - 18, rect.height()),
                         Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft, value)
        painter.restore()


class _BrandsStatusDelegate(ReferenceStatusDelegate):
    def paint(self, painter, option, index):
        item = QStyleOptionViewItem(option)
        self.initStyleOption(item, index)
        text = item.text
        item.text = ""
        style = item.widget.style() if item.widget else QApplication.style()
        style.drawControl(QStyle.ControlElement.CE_ItemViewItem, item, painter, item.widget)
        width = min(option.rect.width() - 8, option.fontMetrics.horizontalAdvance(text) + 32)
        pill = QRectF(option.rect.left() + 25, option.rect.center().y() - 12, width, 24)
        foreground = qc("finance_action_income" if text == "Aktif" else "text_muted")
        outline = QColor(foreground)
        outline.setAlpha(90)
        painter.save()
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setBrush(qc("finance_surface_lowest"))
        painter.setPen(QPen(outline, 1))
        painter.drawRoundedRect(pill, 12, 12)
        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(foreground)
        painter.drawEllipse(QRectF(pill.x() + 8, pill.center().y() - 3, 6, 6))
        painter.setPen(foreground)
        painter.setFont(option.font)
        painter.drawText(pill.adjusted(20, 0, -6, 0), Qt.AlignmentFlag.AlignCenter, text)
        painter.restore()


class BrandsPage(QWidget):
    PAGE_SIZE = 10

    def __init__(self, db, main_window=None):
        super().__init__()
        self.db = db
        self.main_window = main_window
        self.current_page = 0
        self.total_count = 0
        self.init_schema()
        self.init_ui()
        self.load_data()

    def init_schema(self):
        try:
            # Table creation is now handled globally in src/database.py
            if hasattr(self.db, 'create_device_brands_table'):
                self.db.create_device_brands_table()
        except Exception as e:
            from src.utils.logger import logger
            logger.debug(f"BrandsPage local schema init skipped: {e}")

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(26, 54, 26, 20)
        layout.setSpacing(24)

        self.header_grid = QGridLayout()
        title_box = QVBoxLayout()
        self.heading_row = QHBoxLayout()
        self.heading_row.setSpacing(10)
        lbl_title = QLabel("Cihaz Bilgisi & Marka Y\u00f6netimi")
        lbl_title.setWordWrap(True)
        lbl_title.setStyleSheet(theme_qss(
            "font-size: 18px; font-weight: 800; color: @finance_text;"
        ))
        lbl_title.setMaximumWidth(350)
        lbl_title.setWordWrap(False)
        lbl_sub = QLabel(
            "Servis formu ve Teknisyen Panelinde kullan\u0131lacak cihaz t\u00fcr\u00fc ve marka listesini buradan y\u00f6netin."
        )
        lbl_sub.setWordWrap(True)
        lbl_sub.setStyleSheet(theme_qss("font-size: 11px; color: @text_muted;"))
        self.header_badge = QLabel()
        self.header_badge.setStyleSheet(
            "color: #67d8ff; background: #10273a; border: 1px solid #075985; "
            "border-radius: 12px; padding: 4px 10px; font-family: 'JetBrains Mono'; font-weight: 700;"
        )
        self.header_badge.setFixedHeight(32)
        self.heading_row.addWidget(lbl_title)
        self.heading_row.addWidget(self.header_badge)
        self.heading_row.addStretch(1)
        title_box.addLayout(self.heading_row)
        title_box.addWidget(lbl_sub)
        title_box.setSpacing(5)
        self.header_icon = QLabel("\u25a3")
        self.header_icon.setFixedSize(46, 58)
        self.header_icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.header_icon.setFont(QFont("Segoe UI Symbol", 20))
        self.header_icon.setStyleSheet(theme_qss(
            "font-size: 22px; color: @accent; background: @surface_alt; "
            "border: 1px solid @border; border-radius: 10px;"
        ))
        title_box.setContentsMargins(0, 0, 0, 0)
        title_content = QWidget()
        title_content.setLayout(title_box)
        self.header_title = QWidget()
        title_row = QHBoxLayout(self.header_title)
        title_row.setContentsMargins(0, 0, 0, 0)
        title_row.setSpacing(12)
        title_row.addWidget(self.header_icon)
        title_row.addWidget(title_content, 1)
        self.header_title.setMaximumHeight(72)
        self.header_actions = QWidget()
        self.header_actions.setFixedHeight(104)
        self.header_actions.setFixedWidth(654)
        actions = QGridLayout(self.header_actions)
        actions.setContentsMargins(18, 0, 94, 0)
        actions.setSpacing(8)

        btn_catalog = QPushButton("Web Katalo\u011funu G\u00fcncelle")
        btn_catalog.setIcon(_toolbar_icon("catalog"))
        btn_catalog.setIconSize(QSize(18, 18))
        btn_catalog.setFixedSize(250, 48)
        btn_catalog.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_catalog.setStyleSheet(theme_qss(DesignTokens.get_button_qss("secondary")))
        btn_catalog.clicked.connect(self.refresh_reviewed_catalog)
        actions.addWidget(btn_catalog, 0, 0)

        btn_add = QPushButton("Yeni Marka && Cihaz Ekle")
        btn_add.setIcon(_toolbar_icon("add"))
        btn_add.setIconSize(QSize(20, 20))
        btn_add.setFixedSize(238, 48)
        btn_add.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_add.setStyleSheet(theme_qss(DesignTokens.get_button_qss("primary")))
        btn_add.clicked.connect(self.open_add_dialog)
        actions.addWidget(btn_add, 1, 0)

        layout.addLayout(self.header_grid)

        self.filter_frame = QFrame()
        self.filter_frame.setStyleSheet(theme_qss(DesignTokens.get_card_qss()))
        filter_layout = QVBoxLayout(self.filter_frame)
        self.filter_layout = filter_layout
        filter_layout.setContentsMargins(14, 15, 14, 7)
        filter_layout.setAlignment(Qt.AlignmentFlag.AlignTop)
        self.filter_frame.setMinimumHeight(85)
        self.filter_frame.setMaximumHeight(85)
        self.filter_grid = QGridLayout()
        self.filter_grid.setSpacing(10)
        self.filter_grid.setVerticalSpacing(18)
        filter_layout.addLayout(self.filter_grid)
        self.inp_search = QLineEdit()
        self.inp_search.setPlaceholderText("Cihaz t\u00fcr\u00fc veya marka ad\u0131na g\u00f6re filtrele... (Orn: Apple, Tapo, Laptop)")
        search_icon = QPixmap(16, 16)
        search_icon.fill(Qt.GlobalColor.transparent)
        search_painter = QPainter(search_icon)
        search_painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        search_painter.setPen(QPen(QColor("#94a3b8"), 1.6))
        search_painter.drawEllipse(2, 2, 9, 9)
        search_painter.drawLine(10, 10, 14, 14)
        search_painter.end()
        self.inp_search.addAction(QIcon(search_icon), QLineEdit.ActionPosition.LeadingPosition)
        self.inp_search.setMinimumWidth(220)
        self.inp_search.setFixedHeight(40)
        self.inp_search.textChanged.connect(self._reset_and_filter)
        self.cmb_group = ModernComboBox()
        self.cmb_group.set_reference_chevron()
        self.cmb_group.setMinimumWidth(191)
        self.cmb_group.setFixedWidth(191)
        self.cmb_group.setFixedHeight(42)
        self.cmb_group.setStyleSheet(
            self.cmb_group.styleSheet()
            + '\nQComboBox[referenceChevron="true"] { padding-left: 6px; padding-right: 12px; }'
            + '\nQComboBox[referenceChevron="true"]::drop-down { width: 16px; }'
        )
        self.cmb_group.currentTextChanged.connect(self._reset_and_filter)
        self.group_label = QLabel("\u00dcr\u00fcn Grubu:")
        self.group_label.setFixedWidth(84)
        self.cmb_status = ModernComboBox()
        self.cmb_status.set_reference_chevron()
        self.cmb_status.setMinimumWidth(0)
        self.cmb_status.setFixedWidth(124)
        self.cmb_status.setStyleSheet(
            self.cmb_status.styleSheet()
            + '\nQComboBox[referenceChevron="true"] { min-width: 96px; padding-left: 6px; padding-right: 12px; }'
            + '\nQComboBox[referenceChevron="true"]::drop-down { width: 16px; }'
        )
        self.cmb_status.setFixedHeight(42)
        self.cmb_status.addItem("T\u00fcm Durumlar", "all")
        self.cmb_status.addItem("Yaln\u0131zca Aktif", "active")
        self.cmb_status.addItem("Pasif", "inactive")
        self.cmb_status.setCurrentIndex(1)
        self.cmb_status.currentIndexChanged.connect(self._reset_and_filter)
        self.status_label = QLabel("Durum:")
        self.status_label.setFixedWidth(56)
        self.btn_clear_filters = QPushButton("X")
        self.btn_clear_filters.setFixedSize(40, 48)
        self.btn_clear_filters.setToolTip("Filtreleri temizle")
        self.btn_clear_filters.setStyleSheet(
            "QPushButton { color: #b7c5d9; background: #111b2d; border: 1px solid #26364d; "
            "border-radius: 8px; font-size: 15px; font-weight: 700; } "
            "QPushButton:hover { color: #20c5ff; border-color: #20c5ff; }"
        )
        self.btn_clear_filters.clicked.connect(self.clear_filters)
        self.export_buttons = []
        self.export_group = QFrame()
        self.export_group.setStyleSheet(
            "QFrame { background: #111b2d; border: 1px solid #26364d; border-radius: 8px; }"
        )
        self.export_group.setFixedWidth(282)
        export_layout = QHBoxLayout(self.export_group)
        export_layout.setContentsMargins(0, 0, 0, 0)
        export_layout.setSpacing(0)
        for column, (label, callback) in enumerate((
            ("Excel", self.export_excel),
            ("PDF", self.export_pdf),
            ("Yazd\u0131r", self.print_list),
        )):
            button = QPushButton(label)
            button.setIcon(_toolbar_icon(label.casefold() if label != "Yazd\u0131r" else "print"))
            button.setIconSize(QSize(16, 16))
            button.setFixedSize(92, 48)
            button.setStyleSheet(
                "QPushButton { color: #cbd5e1; background: transparent; border: none; "
                "border-right: 1px solid #26364d; padding: 0 8px; font-weight: 600; } "
                "QPushButton:hover { color: #20c5ff; background: #1b2b40; }"
            )
            button.clicked.connect(callback)
            self.export_buttons.append(button)
            export_layout.addWidget(button)
        actions.addWidget(self.export_group, 0, 1, Qt.AlignmentFlag.AlignLeft)
        layout.addWidget(self.filter_frame)
        self._layout_reference_toolbar(self.width())

        self.table_frame = QFrame()
        table_frame = self.table_frame
        table_frame.setStyleSheet(theme_qss(DesignTokens.get_card_qss()))
        tlay = QVBoxLayout(table_frame)
        tlay.setContentsMargins(0, 0, 0, 0)

        self.table = QTableWidget()
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels(
            [
                "",
                "   C\u0130HAZ / \u00dcR\u00dcN GRUBU",
                "   MARKA ADI",
                "\u0130L\u0130\u015eK\u0130L\u0130 SERV\u0130S KAYITLARI",
                "DURUM",
                "\u0130\u015eLEMLER",
            ]
        )
        for header_index in (1, 2, 3, 4):
            self.table.horizontalHeaderItem(header_index).setTextAlignment(
                Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft
            )
        header_view = self.table.horizontalHeader()
        header_view.setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        header_view.setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
        self.table.setColumnWidth(0, 56)
        header_view.setSectionResizeMode(1, QHeaderView.ResizeMode.Fixed)
        header_view.setSectionResizeMode(2, QHeaderView.ResizeMode.Fixed)
        header_view.setSectionResizeMode(3, QHeaderView.ResizeMode.Fixed)
        header_view.setSectionResizeMode(4, QHeaderView.ResizeMode.Fixed)
        header_view.setSectionResizeMode(5, QHeaderView.ResizeMode.Fixed)
        self.table.setColumnWidth(1, 278)
        self.table.setColumnWidth(2, 231)
        self.table.setColumnWidth(3, 290)
        self.table.setColumnWidth(4, 188)
        self.table.setColumnWidth(5, 170)
        header_view.setHighlightSections(False)
        header_view.setFixedHeight(52)
        self.select_all_checkbox = QCheckBox(header_view.viewport())
        self.select_all_checkbox.setFixedSize(18, 18)
        self.select_all_checkbox.move(19, 17)
        self.select_all_checkbox.setStyleSheet(
            "QCheckBox::indicator { width: 16px; height: 16px; border: 1px solid #334155; "
            "border-radius: 4px; background: #0b1220; } "
            "QCheckBox::indicator:checked { background: #0ea5e9; border-color: #38bdf8; }"
        )
        self.select_all_checkbox.stateChanged.connect(self._set_visible_checks)
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionBehavior(
            QAbstractItemView.SelectionBehavior.SelectRows
        )
        self.table.setSelectionMode(
            QAbstractItemView.SelectionMode.SingleSelection
        )
        self.table.setEditTriggers(
            QAbstractItemView.EditTrigger.NoEditTriggers
        )
        self.table.setAlternatingRowColors(False)
        self.table.verticalHeader().setDefaultSectionSize(66)
        self.table.setSortingEnabled(True)
        self.table.setItemDelegateForColumn(4, _BrandsStatusDelegate(self.table))
        self.table.setItemDelegateForColumn(1, _DeviceGroupDelegate(self.table))
        self.table.setStyleSheet(theme_qss(
            DesignTokens.get_table_qss() + "\nQHeaderView::section { font-size: 11px; }"
        ))
        table_font = QFont(self.table.font())
        table_font.setFamily("Inter")
        table_font.setPointSizeF(10.0)
        self.table.setFont(table_font)
        self.table.setMinimumHeight(160)
        tlay.addWidget(self.table, 1)
        table_frame.setMinimumHeight(162)
        layout.addSpacing(8)
        layout.addWidget(table_frame, 0)

        self.footer_grid = QGridLayout()
        self.count_label = QLabel()
        self.count_label.setMinimumWidth(300)
        self.page_size_label = QLabel("Sayfa Ba\u015f\u0131na:")
        self.cmb_page_size = ModernComboBox()
        self.cmb_page_size.set_reference_chevron()
        self.cmb_page_size.setFixedWidth(52)
        self.cmb_page_size.setStyleSheet(theme_qss(
            "QComboBox { min-width: 40px; padding: 2px 14px 2px 4px; } "
            "QComboBox::drop-down { width: 12px; }"
        ))
        for size in (10, 20, 50, 100):
            self.cmb_page_size.addItem(str(size), size)
        self.cmb_page_size.setCurrentIndex(0)
        self.cmb_page_size.currentIndexChanged.connect(self._change_page_size)
        self.btn_previous = QPushButton("\u00d6nceki")
        self.btn_previous.clicked.connect(self.previous_page)
        self.page_label = QLabel()
        self.btn_next = QPushButton("Sonraki")
        self.btn_next.clicked.connect(self.next_page)
        page_button_qss = theme_qss(
            """
            QPushButton {
                min-width: 86px;
                min-height: 32px;
                background: @surface_alt;
                color: @text;
                border: 1px solid @border;
                border-radius: 8px;
                font-weight: 700;
                padding: 4px 10px;
            }
            QPushButton:hover {
                border-color: @accent;
                background: @surface;
            }
            QPushButton:disabled {
                background: @surface;
                color: @text_muted;
                border-color: @border;
            }
            """
        )
        self.btn_previous.setStyleSheet(page_button_qss)
        self.btn_next.setStyleSheet(page_button_qss)
        self.pagination_controls = QWidget()
        navigation = QHBoxLayout(self.pagination_controls)
        navigation.setContentsMargins(0, 0, 0, 0)
        navigation.addWidget(self.btn_previous)
        self.page_buttons_container = QWidget()
        self.page_buttons_layout = QHBoxLayout(self.page_buttons_container)
        self.page_buttons_layout.setContentsMargins(0, 0, 0, 0)
        self.page_buttons_layout.setSpacing(4)
        navigation.addWidget(self.page_buttons_container)
        navigation.addWidget(self.btn_next)
        layout.addLayout(self.footer_grid)
        self._layout_reference_toolbar(self.width())
        self._apply_table_viewport_mode(self.width())

    def _apply_table_viewport_mode(self, page_width):
        if not hasattr(self, 'table_frame'):
            return
        root_layout = self.layout()
        table_index = root_layout.indexOf(self.table_frame)
        if page_width < 1180:
            root_layout.setStretch(table_index, 1)
            table_height = max(220, self.window().height() - 620)
            self.table.setMinimumHeight(0)
            self.table.setMaximumHeight(table_height)
            self.table_frame.setMaximumHeight(table_height + 2)
            compact_widths = (44, 150, 130, 160, 100, 150)
            if page_width < 920:
                compact_widths = (40, 125, 110, 140, 90, 135)
            header_view = self.table.horizontalHeader()
            for column, column_width in enumerate(compact_widths):
                header_view.setSectionResizeMode(column, QHeaderView.ResizeMode.Fixed)
                self.table.setColumnWidth(column, column_width)
        else:
            root_layout.setStretch(table_index, 0)
            self.table.setMinimumHeight(160)
            self.table.setMaximumHeight(16777215)
            self.table_frame.setMaximumHeight(16777215)
        root_layout.invalidate()

    def _layout_reference_toolbar(self, width):
        for grid in (self.header_grid, self.filter_grid):
            while grid.count():
                grid.takeAt(0)
            for column in range(5):
                grid.setColumnStretch(column, 0)
        page_width = min(int(width), self.width())
        narrow = page_width < 1180
        self.layout().setContentsMargins(26, 24 if narrow else 23, 26, 20)
        self.filter_frame.setFixedHeight(122 if narrow else 85)
        self.filter_layout.setContentsMargins(14, 15, 14, 7)
        self.header_grid.addWidget(self.header_title, 0, 0)
        self.header_grid.addWidget(self.header_actions, 1 if narrow else 0, 0 if narrow else 1)
        self.header_grid.setColumnStretch(0, 1)
        if narrow:
            self.filter_grid.addWidget(self.inp_search, 0, 0, 1, 5)
            widgets = (self.group_label, self.cmb_group, self.status_label, self.cmb_status, self.btn_clear_filters)
            for column, widget in enumerate(widgets):
                self.filter_grid.addWidget(widget, 1, column)
            self.filter_grid.setColumnStretch(1, 1)
            self.filter_grid.setColumnStretch(3, 1)
        else:
            widgets = (self.inp_search, self.group_label, self.cmb_group, self.status_label, self.cmb_status, self.btn_clear_filters)
            for column, widget in enumerate(widgets):
                self.filter_grid.addWidget(widget, 0, column)
            self.filter_grid.setColumnStretch(0, 1)
        if hasattr(self, 'footer_grid'):
            while self.footer_grid.count():
                self.footer_grid.takeAt(0)
            self.footer_grid.setAlignment(Qt.AlignmentFlag.AlignTop)
            for column in range(4):
                self.footer_grid.setColumnStretch(column, 0)
            self.footer_grid.addWidget(self.count_label, 0, 0, 1, 3 if narrow else 1)
            self.footer_grid.addWidget(self.page_size_label, 1 if narrow else 0, 0 if narrow else 1)
            self.footer_grid.addWidget(self.cmb_page_size, 1 if narrow else 0, 1 if narrow else 2)
            self.footer_grid.addWidget(self.pagination_controls, 1 if narrow else 0, 2 if narrow else 3)
            if narrow:
                self.footer_grid.setColumnStretch(1, 1)
            else:
                self.footer_grid.setColumnStretch(3, 1)
                self.footer_grid.setAlignment(self.pagination_controls, Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, 'filter_grid'):
            self._layout_reference_toolbar(event.size().width())
            page_width = min(event.size().width(), self.width())
            self._apply_table_viewport_mode(page_width)
            if page_width < 1180:
                compact_widths = (48, 220, 170, 225, 140, 165)
                if page_width < 920:
                    compact_widths = (44, 150, 130, 160, 100, 150)
                header_view = self.table.horizontalHeader()
                for column, column_width in enumerate(compact_widths):
                    header_view.setSectionResizeMode(column, QHeaderView.ResizeMode.Fixed)
                    self.table.setColumnWidth(column, column_width)
            else:
                header_view = self.table.horizontalHeader()
                for column, column_width in ((0, 56), (5, 150)):
                    header_view.setSectionResizeMode(
                        column, QHeaderView.ResizeMode.Fixed
                    )
                    self.table.setColumnWidth(column, column_width)
                for column in (1, 2, 3, 4):
                    header_view.setSectionResizeMode(
                        column, QHeaderView.ResizeMode.Stretch
                    )

            self._fit_table_height()

    def _change_page_size(self):
        self.PAGE_SIZE = int(self.cmb_page_size.currentData())
        self._reset_and_filter()

    def _action_cell(self, button):
        cell = QWidget()
        cell.setMinimumWidth(max(button.minimumWidth(), button.sizeHint().width()) + 12)
        row = QHBoxLayout(cell)
        row.setContentsMargins(6, 0, 6, 0)
        row.addWidget(button, 0, Qt.AlignmentFlag.AlignCenter)
        return cell

    def clear_filters(self):
        self.inp_search.clear()
        self.cmb_group.setCurrentIndex(0)
        self.cmb_status.setCurrentIndex(0)

    @staticmethod
    def _cell_badge(initial, color, size=30):
        badge = QLabel(initial)
        badge.setFixedSize(size, size)
        badge.setAlignment(Qt.AlignmentFlag.AlignCenter)
        badge.setFont(QFont("Segoe UI Symbol", 13 if size > 30 else 9, QFont.Weight.Bold))
        badge.setStyleSheet(
            "QLabel { color: %s; background: #1b2b40; border: 1px solid #34445a; "
            "border-radius: 5px; font-weight: 700; }" % color
        )
        return badge

    @staticmethod
    def _group_glyph(device_type):
        value = device_type.casefold()
        if "ayd\u0131nlatma" in value:
            return "\u2600", "#fbbf24"
        if "ev merkezi" in value or "kamera" in value or "priz" in value:
            return "\u2302", "#a78bfa"
        if "saat" in value:
            return "\u231a", "#c084fc"
        if "telefon" in value:
            return "\u25af", "#34d399"
        if "laptop" in value or "diz\u00fcst\u00fc" in value:
            return "\u25a1", "#38bdf8"
        if "pc" in value or "all in one" in value:
            return "\u25a3", "#22d3ee"
        return "\u25a6", "#22d3ee"

    @staticmethod
    def _brand_initials(brand):
        initials = {
            "tapo": "TP", "aqara": "AQ", "google": "GG", "apple": "AP",
            "samsung": "SM", "appleiphone": "AP", "xiaomi": "MI",
            "asus": "AS", "dell": "DL", "lenovo": "LN",
        }
        return initials.get(brand.casefold().replace(" ", ""), brand[:2].upper())

    def _actions_cell(self, device_type, brand):
        cell = QWidget()
        row = QHBoxLayout(cell)
        row.setContentsMargins(4, 0, 4, 0)
        row.setSpacing(8)
        for label, color, callback in (
            ("D\u00fczenle", "#20c5ff", lambda _checked=False, dt=device_type, br=brand: self.open_edit_dialog(dt, br)),
            ("Sil", "#fb4b73", lambda _checked=False, dt=device_type, br=brand: self.delete_row(dt, br)),
        ):
            button = QPushButton(label)
            button.setCursor(Qt.CursorShape.PointingHandCursor)
            button.setFixedHeight(32)
            button.setFixedWidth(84 if label.startswith("D") else 44)
            button.clicked.connect(callback)
            button.setStyleSheet(
                "QPushButton { color: %s; background: #10192b; border: 1px solid %s; "
                "border-radius: 5px; padding: 0 10px; } "
                "QPushButton:hover { background: #1b2b40; }" % (color, color)
            )
            row.addWidget(button)
        return cell

    def load_data(self):
        try:
            cur = self.db.cursor
            cur.execute(
                "SELECT device_type, brand, is_active FROM device_brands ORDER BY device_type, brand"
            )
            self._all_rows = cur.fetchall() or []
            reference_pairs = (
                ("Ak\u0131ll\u0131 Ayd\u0131nlatma", "Tapo"),
                ("Ak\u0131ll\u0131 Ev Merkezi", "Aqara"),
                ("Ak\u0131ll\u0131 Ev Merkezi", "Google"),
                ("Ak\u0131ll\u0131 Saat", "Apple"),
                ("Ak\u0131ll\u0131 Saat", "Samsung"),
                ("Cep Telefonu", "Apple iPhone"),
                ("Cep Telefonu", "Xiaomi"),
                ("Laptop / Diz\u00fcst\u00fc", "ASUS"),
                ("Laptop / Diz\u00fcst\u00fc", "Dell"),
                ("All In One PC", "Lenovo"),
            )
            order = {pair: index for index, pair in enumerate(reference_pairs)}
            self._all_rows.sort(key=lambda row: (
                0 if (str(self._row_value(row, "device_type", 0)), str(self._row_value(row, "brand", 1))) in order else 1,
                order.get((str(self._row_value(row, "device_type", 0)), str(self._row_value(row, "brand", 1))), 0),
                str(self._row_value(row, "device_type", 0)).casefold(),
                str(self._row_value(row, "brand", 1)).casefold(),
            ))
        except Exception:
            self._all_rows = []
        self.header_badge.setText(f"{len(self._all_rows)} Marka Kay\u0131tl\u0131")
        self._load_service_counts()
        self._populate_group_filter()
        self.apply_filter()

    def _load_service_counts(self):
        self._service_counts = None
        try:
            columns = {row[1] for row in self.db.conn.execute('PRAGMA table_info(devices)')}
            if not {'device_type', 'device_brand'}.issubset(columns):
                return
            where = 'WHERE COALESCE(is_deleted, 0)=0' if 'is_deleted' in columns else ''
            rows = self.db.conn.execute(
                'SELECT device_type, device_brand, COUNT(*) FROM devices '
                + where + ' GROUP BY device_type, device_brand'
            ).fetchall()
            counts = {}
            for device_type, brand, count in rows:
                key = self._service_key(device_type, brand)
                counts[key] = counts.get(key, 0) + int(count)
            self._service_counts = counts
        except Exception:
            # Unavailable counts are not presented as zero business records.
            self._service_counts = None

    @staticmethod
    def _service_key(device_type, brand):
        return (str(device_type or '').strip().casefold(), str(brand or '').strip().casefold())

    def _service_count(self, device_type, brand):
        if self._service_counts is None:
            return None
        return self._service_counts.get(self._service_key(device_type, brand), 0)

    def refresh_reviewed_catalog(self):
        try:
            self.db.refresh_device_model_catalog()
            show_success(
                self,
                "Marka ve model kataloglari guncellendi.",
                "Bilgisayar, cep telefonu ve akilli ev secimleri Yeni Cihaz Ekle alaninda kullanilabilir.",
            )
            self.load_data()
        except Exception as exc:
            show_error(self, f"Katalog guncellenemedi: {exc}")

    def _populate_group_filter(self):
        selected = self.cmb_group.currentData()
        groups = sorted(
            {
                str(self._row_value(row, "device_type", 0) or "").strip()
                for row in self._all_rows
                if str(self._row_value(row, "device_type", 0) or "").strip()
            },
            key=str.casefold,
        )
        self.cmb_group.blockSignals(True)
        self.cmb_group.clear()
        self.cmb_group.addItem(
            f"T\u00fcm \u00dcr\u00fcn Gruplar\u0131 ({len(groups)})", ""
        )
        for group in groups:
            self.cmb_group.addItem(group, group)
        if selected:
            index = self.cmb_group.findData(selected)
            if index >= 0:
                self.cmb_group.setCurrentIndex(index)
        self.cmb_group.blockSignals(False)

    def _reset_and_filter(self):
        self.current_page = 0
        self.apply_filter()

    @staticmethod
    def _row_value(row, key, index):
        if hasattr(row, "keys"):
            return row[key]
        return row[index] if index < len(row) else None

    def _filtered_rows(self):
        text = (self.inp_search.text() or "").strip().casefold()
        selected_group = str(self.cmb_group.currentData() or "").casefold()
        selected_status = str(self.cmb_status.currentData() or "all")
        rows = []
        for row in self._all_rows:
            device_type = str(
                self._row_value(row, "device_type", 0) or ""
            )
            brand = str(self._row_value(row, "brand", 1) or "")
            is_active = bool(int(self._row_value(row, "is_active", 2) or 0))
            if selected_group and device_type.casefold() != selected_group:
                continue
            if selected_status == "active" and not is_active:
                continue
            if selected_status == "inactive" and is_active:
                continue
            if text and text not in f"{device_type} {brand}".casefold():
                continue
            rows.append(row)
        return rows

    def apply_filter(self):
        rows = self._filtered_rows()
        self.total_count = len(rows)
        max_page = max(0, (self.total_count - 1) // self.PAGE_SIZE)
        self.current_page = min(self.current_page, max_page)
        offset = self.current_page * self.PAGE_SIZE
        page_rows = rows[offset:offset + self.PAGE_SIZE]
        self.table.setSortingEnabled(False)
        self.table.setRowCount(0)
        for i, r in enumerate(page_rows):
            self.table.insertRow(i)
            device_type = str(self._row_value(r, "device_type", 0) or "")
            brand = str(self._row_value(r, "brand", 1) or "")
            checkbox = QCheckBox()
            checkbox.setStyleSheet(
                "QCheckBox::indicator { width: 16px; height: 16px; border: 1px solid #334155; "
                "border-radius: 4px; background: #0b1220; } "
                "QCheckBox::indicator:checked { background: #0ea5e9; border-color: #38bdf8; }"
            )
            checkbox_cell = QWidget()
            checkbox_layout = QHBoxLayout(checkbox_cell)
            checkbox_layout.setContentsMargins(0, 0, 0, 0)
            checkbox_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
            checkbox_layout.addWidget(checkbox)
            self.table.setCellWidget(i, 0, checkbox_cell)

            item_type = _ReferenceSortItem("")
            item_type.setData(Qt.ItemDataRole.UserRole, device_type)
            self.table.setItem(i, 1, item_type)

            item_brand = _ReferenceSortItem("")
            item_brand.setData(Qt.ItemDataRole.UserRole, brand)
            self.table.setItem(i, 2, item_brand)
            brand_cell = QWidget()
            brand_layout = QHBoxLayout(brand_cell)
            brand_layout.setContentsMargins(12, 0, 4, 0)
            brand_layout.setSpacing(10)
            brand_layout.addWidget(self._cell_badge(self._brand_initials(brand), "#b9d8ff"))
            brand_label = QLabel(brand)
            brand_label.setStyleSheet(theme_qss(
                "background: transparent; border: 0; color: @text;"
            ))
            brand_layout.addWidget(brand_label, 1)
            self.table.setCellWidget(i, 2, brand_cell)

            count = self._service_count(device_type, brand)
            unit = "Cihaz" if any(term in device_type.casefold() for term in ("ayd\u0131nlatma", "ev merkezi")) else "Fi\u015f"
            item_count = _ServiceCountItem(count, unit)
            item_count.setTextAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
            count_font = QFont(self.table.font())
            count_font.setFamily(DesignTokens.MONO_FONT_NAME)
            item_count.setFont(count_font)
            self.table.setItem(i, 3, item_count)

            active = int(self._row_value(r, "is_active", 2) or 0)
            item_status = QTableWidgetItem("Aktif" if active else "Pasif")
            item_status.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.table.setItem(i, 4, item_status)
            self.table.setCellWidget(i, 5, self._actions_cell(device_type, brand))
        self.table.setSortingEnabled(True)
        start = offset + 1 if self.total_count else 0
        end = min(offset + len(page_rows), self.total_count)
        self.count_label.setText(
            f"Toplam {self.total_count} kay\u0131ttan {start}-{end} "
            "aras\u0131 listeleniyor."
        )
        self.page_label.setText(
            f"{self.current_page + 1} / {max_page + 1}"
        )
        self._refresh_page_buttons(max_page)
        self.btn_previous.setEnabled(self.current_page > 0)
        self.btn_next.setEnabled(self.current_page < max_page)
        self._fit_table_height()

    def _fit_table_height(self):
        if not hasattr(self, "table") or not hasattr(self, "table_frame"):
            return
        natural_height = self.table.horizontalHeader().height() + (
            self.table.rowCount() * self.table.verticalHeader().defaultSectionSize()
        ) + 2
        available_height = max(160, self.height() - (620 if self.width() < 1180 else 315))
        table_height = min(natural_height, available_height)
        self.table.setFixedHeight(table_height)
        self.table_frame.setFixedHeight(table_height + 2)

    def _set_visible_checks(self, state):
        checked = state == int(Qt.CheckState.Checked.value)
        for row in range(self.table.rowCount()):
            cell = self.table.cellWidget(row, 0)
            checkbox = cell.findChild(QCheckBox) if cell else None
            if checkbox:
                checkbox.setChecked(checked)

    def _refresh_page_buttons(self, max_page):
        while self.page_buttons_layout.count():
            item = self.page_buttons_layout.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()
        pages = list(range(min(max_page + 1, 3)))
        if max_page >= 4:
            pages.extend((None, max_page))
        elif max_page == 3:
            pages.append(3)
        for page in pages:
            if page is None:
                self.page_buttons_layout.addWidget(QLabel("..."))
                continue
            button = QPushButton(str(page + 1))
            button.setFixedSize(34, 32)
            button.setCursor(Qt.CursorShape.PointingHandCursor)
            active = page == self.current_page
            button.setStyleSheet(
                "QPushButton { color: %s; background: %s; border: 1px solid %s; "
                "border-radius: 7px; font-weight: 700; }"
                % (("#06111f", "#19bff5", "#19bff5") if active else
                   ("#cbd5e1", "#111b2d", "#26364d"))
            )
            button.clicked.connect(
                lambda _checked=False, value=page: self._go_to_page(value)
            )
            self.page_buttons_layout.addWidget(button)

    def _go_to_page(self, page):
        self.current_page = int(page)
        self.apply_filter()

    def open_add_dialog(self):
        from src.ui.dialogs.brand_edit_dialog import BrandEditDialog

        dlg = BrandEditDialog(self.db, self)
        if dlg.exec():
            show_success(self, "Marka kaydedildi.")
            self.load_data()

    def open_edit_dialog(self, device_type, brand):
        from src.ui.dialogs.brand_edit_dialog import BrandEditDialog

        dialog = BrandEditDialog(
            self.db,
            self,
            initial_device_type=device_type,
            initial_brand=brand,
            edit_mode=True,
        )
        if dialog.exec():
            show_success(self, "Marka g\u00fcncellendi.")
            self.load_data()

    def previous_page(self):
        if self.current_page > 0:
            self.current_page -= 1
            self.apply_filter()

    def next_page(self):
        if (self.current_page + 1) * self.PAGE_SIZE < self.total_count:
            self.current_page += 1
            self.apply_filter()

    def _output_rows(self):
        return [
            (
                str(self._row_value(row, "device_type", 0) or ""),
                str(self._row_value(row, "brand", 1) or ""),
                (
                    "Aktif"
                    if int(self._row_value(row, "is_active", 2) or 0)
                    else "Pasif"
                ),
            )
            for row in self._filtered_rows()
        ]

    def _export_rows(self):
        rows = []
        for device_type, brand, status in self._output_rows():
            count = self._service_count(device_type, brand)
            rows.append((device_type, brand, '-' if count is None else count, status))
        return rows

    def export_excel(self):
        path, _selected = QFileDialog.getSaveFileName(
            self,
            "Excel \u00c7\u0131kt\u0131",
            "markalar.xlsx",
            "Excel (*.xlsx)",
        )
        if not path:
            return
        try:
            from openpyxl import Workbook

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Markalar"
            sheet.append(["Cihaz Turu", "Marka", "Iliskili Servis Kayitlari", "Durum"])
            for row in self._export_rows():
                sheet.append(list(row))
            workbook.save(path)
            show_success(self, f"Excel kaydedildi: {path}")
        except Exception as exc:
            show_error(self, f"Excel olusturulamadi: {exc}")

    def _html_document(self):
        rows = "".join(
            "<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
                html.escape(device_type),
                html.escape(brand),
                html.escape(str(count)),
                html.escape(status),
            )
            for device_type, brand, count, status in self._export_rows()
        )
        return (
            "<h2>Marka Adlari Yonetimi</h2>"
            "<table border='1' cellspacing='0' cellpadding='6' "
            "width='100%'><tr><th>Cihaz Turu</th><th>Marka</th>"
            f"<th>Iliskili Servis Kayitlari</th><th>Durum</th></tr>{rows}</table>"
        )

    def export_pdf(self):
        path, _selected = QFileDialog.getSaveFileName(
            self,
            "PDF Kaydet",
            "markalar.pdf",
            "PDF (*.pdf)",
        )
        if not path:
            return
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(str(Path(path)))
        document = QTextDocument()
        document.setHtml(self._html_document())
        document.print(printer)
        show_success(self, f"PDF kaydedildi: {path}")

    def print_list(self):
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        dialog = QPrintDialog(printer, self)
        if not dialog.exec():
            return
        document = QTextDocument()
        document.setHtml(self._html_document())
        document.print(printer)

    def delete_row(self, device_type, brand):
        try:
            self.db.cursor.execute(
                "DELETE FROM device_brands WHERE device_type=? AND brand=?",
                (device_type, brand),
            )
            self.db.conn.commit()
            show_success(self, "Kayıt silindi.")
            self.load_data()
        except Exception as e:
            show_error(self, f"Silme hatası: {e}")

