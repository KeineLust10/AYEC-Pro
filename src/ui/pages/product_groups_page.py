import html
from datetime import datetime, timedelta
from pathlib import Path

from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor, QPalette, QTextDocument
from PyQt6.QtPrintSupport import QPrintDialog, QPrinter
from PyQt6.QtWidgets import (
    QAbstractItemView,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QGridLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QComboBox,
    QLayout,
    QPushButton,
    QScrollArea,
    QStackedWidget,
    QTabBar,
    QTableWidget,
    QTableWidgetItem,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
)

from src.ui.dialogs.modern_input_dialog import ModernInputDialog
from src.ui.widgets.modern_dialog import ModernDialog
from src.ui.dialogs.simple_confirm import SimpleConfirmDialog
from src.utils.design_system import DesignTokens
from src.ui.widgets.reference_status_delegate import ReferenceStatusDelegate
from src.ui.widgets.navigation_glyph import NavigationGlyph
from src.utils.theme_colors import theme_qss, tc
from src.utils.toast_notification import (
    show_error,
    show_info,
    show_success,
    show_warning,
)


class ProductGroupsPage(QWidget):
    PAGE_SIZE = 10
    STANDARD_TEMPLATES = (
        (("all in one", "aio"), "Entegre Sistem \u015eablonu (AIO-01)"),
        (("bilgisayar",), "Donan\u0131m & Test \u015eablonu (D-02)"),
        (("laptop", "pc"), "Bili\u015fim & Donan\u0131m \u015eablonu (F-201)"),
        (("elektronik",), "Elektronik Devre \u015eablonu (F-104)"),
        (("telefon", "mobil"), "Mobil Cihaz Kontrol \u015eablonu (M-01)"),
        (("saat", "giyilebilir"), "Giyilebilir Donan\u0131m Testi (W-03)"),
        (("priz", "iot"), "IoT Voltaj ve Yaz\u0131l\u0131m \u015eablonu"),
        (("kamera", "cctv", "g\u00fcvenlik"), "Optik & A\u011f G\u00fcvenlik Testi (CAM-2)"),
        (("merkez", "hub"), "Merkezi Hub Haberle\u015fme Protokol\u00fc"),
        (("ayd\u0131nlatma", "led"), "LED S\u00fcr\u00fcc\u00fc ve Renk Kalibrasyonu"),
    )
    GROUP_DESCRIPTIONS = (
        (("laptop",), "Diz\u00fcst\u00fc Bilgisayar & Ultrabook"),
        (("elektronik",), "Genel Elektronik & Kart Onar\u0131m\u0131"),
        (("telefon",), "Ak\u0131ll\u0131 Telefon & Tablet Cihazlar"),
        (("bilgisayar", "pc"), "Masa\u00fcst\u00fc, Kasa & \u0130\u015f \u0130stasyonlar\u0131"),
        (("all in one", "aio"), "Entegre Monit\u00f6r Bilgisayarlar"),
        (("saat",), "Giyilebilir Teknoloji & Bileklik"),
        (("priz",), "IoT Enerji & R\u00f6le \u00dcniteleri"),
        (("kamera", "cctv"), "IP Kamera, CCTV ve Sens\u00f6rler"),
        (("merkez", "hub"), "Gateway, Zigbee/Z-Wave Hub"),
        (("ayd\u0131nlatma",), "RGB Panel, LED \u015eerit & Armat\u00fcr"),
    )

    def __init__(self, db, main_window=None):
        super().__init__()
        self.db = db
        self.main_window = main_window
        self.current_page = 0
        self.total_count = 0
        self._ensure_schema()
        self._seed_existing_groups()
        self._build_ui()
        self.load_data()
        self._apply_responsive_layout()

    def _ensure_schema(self):
        self.db.cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS product_groups (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL COLLATE NOCASE UNIQUE,
                is_active INTEGER DEFAULT 1,
                created_at TEXT,
                updated_at TEXT
            )
            """
        )
        self.db.conn.commit()

    def _seed_existing_groups(self):
        names = set()
        for query in (
            """
            SELECT DISTINCT device_type
            FROM device_brands
            WHERE TRIM(COALESCE(device_type, '')) <> ''
            """,
            """
            SELECT DISTINCT device_type
            FROM devices
            WHERE TRIM(COALESCE(device_type, '')) <> ''
            """,
        ):
            try:
                self.db.cursor.execute(query)
                names.update(
                    str(row[0]).strip()
                    for row in (self.db.cursor.fetchall() or [])
                    if row[0]
                )
            except Exception:
                continue
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        for name in sorted(names):
            self.db.cursor.execute(
                """
                INSERT OR IGNORE INTO product_groups (
                    name, is_active, created_at, updated_at
                )
                VALUES (?, 1, ?, ?)
                """,
                (name, now, now),
            )
        self.db.conn.commit()

    def _build_ui(self):
        self.setMinimumSize(0, 0)
        root = QVBoxLayout(self)
        self._root_layout = root
        root.setSizeConstraint(QLayout.SizeConstraint.SetNoConstraint)
        root.setContentsMargins(24, 24, 24, 20)
        root.setSpacing(13)

        header = QHBoxLayout()
        title_box = QVBoxLayout()
        title = QLabel("\u00dcr\u00fcn Grubu Y\u00f6netimi")
        title.setStyleSheet(
            theme_qss("font-size: 18px; font-weight: 800; color: @text; background: transparent; border: none;")
        )
        subtitle = QLabel(
            "Servis kay\u0131tlar\u0131nda kullan\u0131lacak cihaz ve "
            "\u00fcr\u00fcn gruplar\u0131n\u0131 y\u00f6netin."
        )
        subtitle.setStyleSheet(
            theme_qss("font-size: 10px; color: @text_muted; background: transparent; border: none;")
        )
        title_row = QHBoxLayout()
        title_row.setSpacing(10)
        title_row.addWidget(title)
        self.active_category_badge = QLabel("0 Aktif Kategori")
        self.active_category_badge.setStyleSheet(theme_qss(
            "QLabel { color: @accent; background: @surface_alt; border: 1px solid @accent; "
            "border-radius: 10px; padding: 3px 8px; font-size: 9px; font-weight: 700; }"
        ))
        self.active_category_badge.setFixedHeight(22)
        title_row.addWidget(self.active_category_badge, 0, Qt.AlignmentFlag.AlignVCenter)
        title_box.addLayout(title_row)
        title_box.addWidget(subtitle)
        header.addLayout(title_box)
        header.addStretch()
        btn_hierarchy = QPushButton("Grup Hiyerar\u015fisi")
        btn_hierarchy.setFixedSize(100, 26)
        btn_hierarchy.setStyleSheet(theme_qss(
            DesignTokens.get_button_qss("secondary") +
            " QPushButton { font-size: 9px; padding: 0 6px; }"
        ))
        btn_hierarchy.clicked.connect(self.show_group_hierarchy)
        header.addWidget(btn_hierarchy)
        btn_add = QPushButton("Yeni \u00dcr\u00fcn Grubu Ekle")
        btn_add.setFixedSize(128, 26)
        btn_add.setStyleSheet(
            theme_qss(DesignTokens.get_button_qss("primary") +
                      " QPushButton { font-size: 9px; padding: 0 6px; }")
        )
        btn_add.clicked.connect(self.add_group)
        header.addWidget(btn_add)
        root.addLayout(header)
        root.addSpacing(6)

        metrics = QHBoxLayout()
        metrics.setSpacing(10)
        self.metric_total_groups = self._metric_card("Toplam \u00dcr\u00fcn Grubu", "0")
        self.metric_brand_pool = self._metric_card("E\u015fle\u015fen Marka Havuzu", "0")
        self.metric_device_inventory = self._metric_card("Kay\u0131tl\u0131 Cihaz Envanteri", "-")
        self.metric_service_volume = self._metric_card("En Y\u00fcksek Servis Hacmi", "-")
        self.active_group_summary = QLabel()
        self.active_group_summary.setStyleSheet(theme_qss('color: @finance_action_income; font-size: 9px;'))
        self.metric_total_groups.detail_layout.addWidget(self.active_group_summary)
        for metric in (
            self.metric_total_groups,
            self.metric_device_inventory,
            self.metric_service_volume,
            self.metric_brand_pool,
        ):
            metrics.addWidget(metric, 1)
        root.addLayout(metrics)

        toolbar_frame = QFrame()
        toolbar_frame.setMinimumWidth(0)
        toolbar_frame.setStyleSheet(theme_qss(
            'QFrame { background: @surface; border: 1px solid @border; border-radius: 8px; }'
        ))
        toolbar_inner = QHBoxLayout(toolbar_frame)
        toolbar_inner.setContentsMargins(12, 8, 12, 8)
        toolbar_inner.setSpacing(8)
        toolbar = QHBoxLayout()
        self.search = QLineEdit()
        self.search.setObjectName("ProductGroupSearch")
        self.search.setPlaceholderText(
            "\u00dcr\u00fcn grubu veya kategori ara (\u00f6rn: Laptop, Priz, Ak\u0131ll\u0131 Ev)..."
        )
        self.search.setStyleSheet(theme_qss(
            "QLineEdit#ProductGroupSearch { background: @surface_alt; color: @text; "
            "border: 1px solid @border; border-radius: 5px; font-size: 9px; padding: 0 8px; }"
        ))
        search_palette = self.search.palette()
        search_palette.setColor(QPalette.ColorRole.Base, QColor(tc("surface_alt")))
        search_palette.setColor(QPalette.ColorRole.Text, QColor(tc("text")))
        search_palette.setColor(QPalette.ColorRole.PlaceholderText, QColor(tc("text_muted")))
        self.search.setPalette(search_palette)
        self.search.textChanged.connect(self._reset_and_load)
        self.search.setFixedWidth(266)
        self.search.setFixedHeight(22)
        self.search_frame = QFrame()
        self.search_frame.setObjectName("ProductGroupSearchFrame")
        self.search_frame.setFixedSize(286, 26)
        self.search_frame.setStyleSheet(theme_qss(
            "QFrame#ProductGroupSearchFrame { background: @surface_alt; "
            "border: 1px solid @border; border-radius: 5px; }"
        ))
        search_frame_layout = QHBoxLayout(self.search_frame)
        search_frame_layout.setContentsMargins(8, 1, 8, 1)
        search_frame_layout.setSpacing(4)
        search_icon = NavigationGlyph("\U0001f50d", title="Search")
        search_icon.setFixedSize(14, 14)
        search_icon.set_render_color(tc("text_muted"))
        search_icon.set_navigation_state(True)
        search_frame_layout.addWidget(search_icon)
        self.search.setStyleSheet(theme_qss(
            "QLineEdit#ProductGroupSearch { background: transparent; color: @text; "
            "border: none; font-size: 9px; padding: 0; }"
        ))
        search_frame_layout.addWidget(self.search, 1)
        toolbar.addWidget(self.search_frame)
        self.view_tabs = QTabBar()
        self.view_tabs.setDrawBase(False)
        self.view_tabs.addTab('Liste')
        self.view_tabs.addTab('Kart')
        self.view_tabs.setFixedWidth(96)
        self.view_tabs.setStyleSheet(theme_qss(
            'QTabBar::tab { color: @text_muted; background: @finance_surface_lowest; '
            'border: 1px solid @finance_outline; padding: 6px 8px; font-size: 9px; } '
            'QTabBar::tab:selected { background: @finance_active_bg; color: @finance_primary; }'
        ))
        toolbar.addStretch(1)
        toolbar.addWidget(self.view_tabs)
        for label, callback in (
            ("Excel \u00c7\u0131kt\u0131", self.export_excel),
            ("PDF Kaydet", self.export_pdf),
            ("Yazd\u0131r", self.print_list),
        ):
            button = QPushButton(label)
            button.setStyleSheet(
                theme_qss(DesignTokens.get_button_qss("secondary"))
            )
            button.setFixedHeight(32)
            button.setFixedWidth({"Excel \u00c7\u0131kt\u0131": 72, "PDF Kaydet": 78, "Yazd\u0131r": 56}[label])
            button.setStyleSheet(theme_qss(
                DesignTokens.get_button_qss("secondary") + " QPushButton { font-size: 9px; padding: 0 4px; }"
            ))
            button.clicked.connect(callback)
            toolbar.addWidget(button)
        root.addSpacing(5)
        toolbar_inner.addLayout(toolbar)
        root.addWidget(toolbar_frame)
        root.addSpacing(3)

        card = QFrame()
        self.table_card = card
        card.setMinimumWidth(0)
        card.setStyleSheet(
            theme_qss(DesignTokens.get_card_qss(hover=False))
        )
        card_layout = QVBoxLayout(card)
        card_layout.setContentsMargins(0, 0, 0, 0)
        card_layout.setSpacing(2)
        self.table = QTableWidget(0, 6)
        self.table.setMinimumWidth(0)
        self.table.setHorizontalHeaderLabels(
            [
                "#",
                "\u00dcR\u00dcN GRUBU ADI",
                "BA\u011eLI MARKA / KATEGOR\u0130",
                "STANDART SERV\u0130S \u015eABLONU",
                "DURUM",
                "AKS\u0130YONLAR",
            ]
        )
        header_view = self.table.horizontalHeader()
        header_view.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        for section, width in ((0, 48), (1, 233), (2, 256), (3, 205), (4, 88), (5, 128)):
            self.table.setColumnWidth(section, width)
        for section in (1, 2, 3):
            header_view.setSectionResizeMode(section, QHeaderView.ResizeMode.Stretch)
        header_view.setDefaultAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
        header_view.setSortIndicatorShown(False)
        number_header = self.table.horizontalHeaderItem(0)
        number_header.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        self.table.setSelectionBehavior(
            QAbstractItemView.SelectionBehavior.SelectRows
        )
        self.table.setEditTriggers(
            QAbstractItemView.EditTrigger.NoEditTriggers
        )
        self.table.setSortingEnabled(True)
        self.table.sortItems(0, Qt.SortOrder.AscendingOrder)
        header_view.setSortIndicatorShown(False)
        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setDefaultSectionSize(49)
        self.table.setItemDelegateForColumn(4, ReferenceStatusDelegate(self.table))
        self.table.setStyleSheet(theme_qss(
            DesignTokens.get_table_qss() +
            ' QTableWidget { font-size: 11px; } '
            'QHeaderView::section { font-size: 11px; padding: 5px 7px; }'
        ))
        self.table.horizontalHeader().setFixedHeight(29)
        self.table.horizontalHeader().setStretchLastSection(False)
        card_layout.addWidget(self.table)
        footer = QHBoxLayout()
        self._table_footer = footer
        footer.setContentsMargins(12, 7, 12, 7)
        self.count_label = QLabel()
        self.count_label.setStyleSheet(theme_qss(
            'QLabel { color: @text_muted; font-size: 9px; background: transparent; border: none; }'
        ))
        footer.addWidget(self.count_label)
        footer.addWidget(QLabel("|"))
        page_size_label = QLabel("Sayfa Ba\u015f\u0131na:")
        page_size_label.setStyleSheet(theme_qss('color: @text_muted; font-size: 9px;'))
        footer.addWidget(page_size_label)
        self.page_size_combo = QComboBox()
        self.page_size_combo.addItems(["10", "25", "50"])
        self.page_size_combo.setCurrentText(str(self.PAGE_SIZE))
        self.page_size_combo.setFixedSize(48, 22)
        self.page_size_combo.setStyleSheet(theme_qss(
            'QComboBox { background: @surface; color: @text; border: 1px solid @border; '
            'border-radius: 4px; padding: 0 4px; font-size: 9px; } '
            'QComboBox::drop-down { width: 16px; border-left: 1px solid @border; }'
        ))
        self.page_size_combo.currentTextChanged.connect(self._change_page_size)
        footer.addWidget(self.page_size_combo)
        footer.addStretch()
        self.btn_previous = QPushButton("\u00d6nceki")
        self.btn_previous.clicked.connect(self.previous_page)
        self.page_label = QLabel()
        self.page_label.setStyleSheet(theme_qss(
            'QLabel { color: @text; font-family: "JetBrains Mono"; font-size: 9px; background: transparent; border: none; }'
        ))
        self.btn_next = QPushButton("Sonraki")
        self.btn_next.clicked.connect(self.next_page)
        page_button_qss = theme_qss(
            """
            QPushButton {
                min-width: 0px;
                min-height: 20px;
                background: @surface_alt;
                color: @text;
                border: 1px solid @border;
                border-radius: 8px;
                font-weight: 700;
                padding: 0 3px;
            }
            QPushButton:hover { border-color: @accent; background: @surface; }
            QPushButton:disabled { background: @surface; color: @text_muted; border-color: @border; }
            """
        )
        self.btn_previous.setStyleSheet(page_button_qss)
        self.btn_next.setStyleSheet(page_button_qss)
        self.btn_previous.setFixedSize(58, 24)
        self.btn_next.setFixedSize(58, 24)
        footer.addWidget(self.btn_previous)
        self.page_buttons_host = QWidget()
        self.page_buttons_layout = QHBoxLayout(self.page_buttons_host)
        self.page_buttons_layout.setContentsMargins(0, 0, 0, 0)
        self.page_buttons_layout.setSpacing(4)
        footer.addWidget(self.page_buttons_host)
        footer.addWidget(self.btn_next)
        card_layout.addLayout(footer)
        self.views = QStackedWidget()
        self.views.addWidget(card)
        self.cards_scroll = QScrollArea()
        self.cards_scroll.setWidgetResizable(True)
        self.cards_scroll.setFrameShape(QFrame.Shape.NoFrame)
        self.cards_host = QWidget()
        self.cards_grid = QGridLayout(self.cards_host)
        self.cards_grid.setContentsMargins(0, 0, 0, 0)
        self.cards_grid.setSpacing(12)
        self.cards_grid.setAlignment(Qt.AlignmentFlag.AlignTop)
        self.cards_empty = QLabel('Bu arama i\u00e7in \u00fcr\u00fcn grubu bulunamad\u0131.', self.cards_host)
        self.cards_empty.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.cards_empty.setStyleSheet(theme_qss('color: @text_muted; padding: 32px;'))
        self._group_cards = []
        self.cards_scroll.setWidget(self.cards_host)
        self.views.addWidget(self.cards_scroll)
        self.view_tabs.currentChanged.connect(self.views.setCurrentIndex)
        root.addWidget(self.views, 1)


    def _metric_card(self, title, value):
        card = QFrame()
        card.setObjectName("ProductGroupMetric")
        card.setFixedHeight(82)
        card.setStyleSheet(theme_qss(
            "QFrame#ProductGroupMetric { background: @surface; border: 1px solid @border; "
            "border-radius: 8px; padding: 8px; }"
        ))
        layout = QHBoxLayout(card)
        layout.setContentsMargins(12, 8, 12, 8)
        layout.setSpacing(8)
        detail_layout = QVBoxLayout()
        detail_layout.setContentsMargins(0, 0, 0, 0)
        detail_layout.setSpacing(2)
        label = QLabel(title)
        label.setStyleSheet(theme_qss("color: @text_muted; font-size: 10px; font-weight: 700; background: transparent; border: none;"))
        number = QLabel(value)
        number.setWordWrap(False)
        number.setObjectName("metricValue")
        number.setStyleSheet(theme_qss("color: @text; font-size: 16px; font-weight: 800; background: transparent; border: none;"))
        detail_layout.addWidget(label)
        value_row = QHBoxLayout()
        value_row.setContentsMargins(0, 0, 0, 0)
        value_row.setSpacing(4)
        value_row.addWidget(number)
        unit = {
            "Toplam \u00dcr\u00fcn Grubu": "Grup",
            "Kay\u0131tl\u0131 Cihaz Envanteri": "Adet",
            "E\u015fle\u015fen Marka Havuzu": "Marka",
        }.get(title)
        if unit:
            unit_label = QLabel(unit)
            unit_label.setStyleSheet(theme_qss('color: @text_muted; font-size: 9px; background: transparent; border: none;'))
            value_row.addWidget(unit_label, 0, Qt.AlignmentFlag.AlignBottom)
        value_row.addStretch(1)
        detail_layout.addLayout(value_row)
        detail = QLabel()
        detail.setStyleSheet(theme_qss('color: @text_muted; font-size: 9px; background: transparent; border: none;'))
        detail_layout.addWidget(detail)
        layout.addLayout(detail_layout, 1)
        icon_glyph, icon_color, icon_token = {
            "Toplam \u00dcr\u00fcn Grubu": ("\U0001f4e6", "#a78bfa", "finance_primary"),
            "Kay\u0131tl\u0131 Cihaz Envanteri": ("\U0001f4f1", "#22d3ee", "finance_action_income"),
            "En Y\u00fcksek Servis Hacmi": ("\U0001f4bb", "@warning", "warning"),
            "E\u015fle\u015fen Marka Havuzu": ("\U0001f9e9", "#c084fc", "accent"),
        }.get(title, ("\U0001f4e6", "@accent", "accent"))
        icon = NavigationGlyph(icon_glyph, title=title, active_token=icon_token)
        icon.setFixedSize(34, 34)
        icon.set_navigation_state(True)
        icon.set_render_color(icon_color if icon_color.startswith("#") else "")
        icon.setStyleSheet(theme_qss(
            "QLabel { color: %s; background: @surface_alt; border: 1px solid %s; "
            "border-radius: 8px; font-size: 16px; }" % (icon_color, icon_color)
        ))
        layout.addWidget(icon, 0, Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        card.metric_value = number
        card.detail_layout = detail_layout
        card.metric_detail = detail
        return card

    def _reset_and_load(self):
        self.current_page = 0
        self.load_data()

    def _where_clause(self):
        query = self.search.text().strip()
        if not query:
            return "", []
        return "WHERE name LIKE ?", [f"%{query}%"]

    def load_data(self):
        self._refresh_metrics()
        self._load_linked_brands()
        where_sql, params = self._where_clause()
        self.db.cursor.execute(
            f"SELECT COUNT(*) FROM product_groups {where_sql}",
            tuple(params),
        )
        self.total_count = int(self.db.cursor.fetchone()[0] or 0)
        max_page = max(0, (self.total_count - 1) // self.PAGE_SIZE)
        self.current_page = min(self.current_page, max_page)
        offset = self.current_page * self.PAGE_SIZE
        self.db.cursor.execute(
            f"""
            SELECT id, name, is_active
            FROM product_groups
            {where_sql}
            ORDER BY id ASC
            LIMIT ? OFFSET ?
            """,
            tuple(params + [self.PAGE_SIZE, offset]),
        )
        rows = self.db.cursor.fetchall() or []
        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(rows))
        for row_index, row in enumerate(rows):
            group_id = int(row[0])
            name = str(row[1] or "")
            active = bool(row[2])
            name_item = QTableWidgetItem(name)
            name_item.setData(Qt.ItemDataRole.UserRole, group_id)
            self.table.setItem(row_index, 1, name_item)
            self.table.setCellWidget(row_index, 1, self._group_name_cell(name))
            number_item = QTableWidgetItem(f"{offset + row_index + 1:02d}")
            number_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.table.setItem(row_index, 0, number_item)
            brands = self._linked_brands_for(name)
            summary = '-' if brands is None else '{} Marka'.format(len(brands))
            if brands:
                summary += ' ({})'.format(', '.join(brands[:3]) + (', ...' if len(brands) > 3 else ''))
            brands_item = QTableWidgetItem(summary)
            brands_item.setToolTip('Aktif markalar: ' + ', '.join(brands or []))
            self.table.setItem(row_index, 2, brands_item)
            self.table.setCellWidget(row_index, 2, self._brands_cell(summary))
            template_item = QTableWidgetItem(self._standard_template(name))
            template_item.setToolTip(template_item.text())
            self.table.setItem(row_index, 3, template_item)
            self.table.setCellWidget(row_index, 3, self._template_cell(template_item.text()))
            self.table.setItem(
                row_index,
                4,
                QTableWidgetItem(
                    "Aktif" if active else "Pasif"
                ),
            )
            btn_edit = QPushButton("D\u00fczenle")
            btn_edit.setCursor(Qt.CursorShape.PointingHandCursor)
            btn_edit.setMinimumWidth(66)
            btn_edit.setFixedSize(66, 26)
            btn_edit.setStyleSheet(
                theme_qss(
                    'QPushButton { background: @finance_surface_lowest; color: @finance_primary; '
                    'border: 1px solid @finance_outline; border-radius: 4px; padding: 0 2px; font-size: 10px; } '
                    'QPushButton:hover { border-color: @finance_primary; }'
                )
            )
            btn_edit.clicked.connect(
                lambda _checked=False, gid=group_id, value=name: (
                    self.edit_group(gid, value)
                )
            )
            btn_delete = QPushButton("Sil")
            btn_delete.setCursor(Qt.CursorShape.PointingHandCursor)
            btn_delete.setMinimumWidth(34)
            btn_delete.setFixedSize(34, 26)
            btn_delete.setStyleSheet(
                theme_qss(
                    'QPushButton { background: @finance_surface_lowest; color: @finance_action_expense; '
                    'border: 1px solid @finance_action_expense; border-radius: 4px; padding: 0 2px; font-size: 10px; } '
                    'QPushButton:hover { background: @finance_surface_high; }'
                )
            )
            btn_delete.clicked.connect(
                lambda _checked=False, gid=group_id, value=name: (
                    self.delete_group(gid, value)
                )
            )
            self.table.setCellWidget(row_index, 5, self._actions_cell(btn_edit, btn_delete))
        self.table.setSortingEnabled(True)
        self._render_group_cards(rows)
        start = offset + 1 if self.total_count else 0
        end = min(offset + len(rows), self.total_count)
        self.count_label.setText(
            f"Toplam {self.total_count} kay\u0131ttan {start}-{end} "
            "aras\u0131 g\u00f6steriliyor."
        )
        self.page_label.setText(f"{self.current_page + 1} / {max_page + 1}")
        self.page_label.setVisible(False)
        self._render_page_buttons(max_page + 1)
        self.btn_previous.setEnabled(self.current_page > 0)
        self.btn_next.setEnabled(self.current_page < max_page)
        self._fit_table_card()

    def _fit_table_card(self):
        if not hasattr(self, "table_card") or not hasattr(self, "table"):
            return
        natural_table_height = (
            self.table.horizontalHeader().height()
            + self.table.rowCount() * self.table.verticalHeader().defaultSectionSize()
        )
        footer_height = self._table_footer.sizeHint().height()
        natural_card_height = natural_table_height + footer_height + 4
        available_card_height = max(
            220, self.height() - (344 if self.width() < 1200 else 256)
        )
        card_height = min(natural_card_height, available_card_height)
        table_height = max(140, card_height - footer_height - 4)
        self.table.setFixedHeight(table_height)
        self.table_card.setFixedHeight(card_height)

    def _render_page_buttons(self, page_count):
        while self.page_buttons_layout.count():
            item = self.page_buttons_layout.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()
        if page_count > 8:
            label = QLabel(self.page_label.text())
            label.setStyleSheet(theme_qss('color: @text_muted; font-size: 9px;'))
            self.page_buttons_layout.addWidget(label)
            return
        for page_index in range(page_count):
            button = QPushButton(str(page_index + 1))
            button.setFixedSize(24, 26)
            button.setStyleSheet(theme_qss(
                'QPushButton { background: %s; color: %s; border: 1px solid %s; '
                'border-radius: 5px; font-size: 9px; padding: 0; }' % (
                    '@accent' if page_index == self.current_page else '@surface_alt',
                    '@accent_foreground' if page_index == self.current_page else '@text_muted',
                    '@accent' if page_index == self.current_page else '@border',
                )
            ))
            button.clicked.connect(lambda _checked=False, index=page_index: self.go_to_page(index))
            self.page_buttons_layout.addWidget(button)

    def go_to_page(self, page_index):
        self.current_page = max(0, int(page_index))
        self.load_data()

    def _change_page_size(self, value):
        if not str(value).isdigit():
            return
        self.PAGE_SIZE = int(value)
        self.current_page = 0
        self.load_data()

    def _render_group_cards(self, rows):
        while self.cards_grid.count():
            item = self.cards_grid.takeAt(0)
            if item.widget() and item.widget() is not self.cards_empty:
                item.widget().deleteLater()
        self._group_cards = []
        for group_id, name, active in rows:
            name = str(name or '')
            card = QFrame()
            card.setObjectName('ProductGroupReferenceCard')
            card.setStyleSheet(theme_qss(
                'QFrame#ProductGroupReferenceCard { background: @finance_surface; '
                'border: 1px solid @finance_outline; border-radius: 8px; }'
            ))
            body = QVBoxLayout(card)
            body.setContentsMargins(16, 16, 16, 16)
            title = QLabel(name)
            title.setWordWrap(True)
            title.setStyleSheet(theme_qss('font-size: 16px; font-weight: 700; color: @finance_text;'))
            body.addWidget(title)
            brands = self._linked_brands_for(name)
            label = QLabel('-' if brands is None else '{} Marka: {}'.format(len(brands), ', '.join(brands)))
            label.setWordWrap(True)
            label.setStyleSheet(theme_qss('color: @text_muted; font-size: 12px;'))
            body.addWidget(label)
            status = QLabel('Aktif' if active else 'Pasif')
            status.setStyleSheet(theme_qss('color: @finance_action_income;' if active else 'color: @text_muted;'))
            body.addWidget(status)
            actions = QHBoxLayout()
            for text, callback, variant in (
                ('D\u00fczenle', self.edit_group, 'secondary'),
                ('Sil', self.delete_group, 'destructive'),
            ):
                button = QPushButton(text)
                color = '@finance_action_expense' if variant == 'destructive' else '@finance_primary'
                button.setMinimumHeight(32)
                button.setStyleSheet(theme_qss(
                    'QPushButton { background: @finance_surface_lowest; color: ' + color + '; '
                    'border: 1px solid ' + color + '; border-radius: 4px; padding: 0 8px; } '
                    'QPushButton:hover { background: @finance_surface_high; }'
                ))
                button.clicked.connect(lambda _checked=False, gid=int(group_id), value=name, fn=callback: fn(gid, value))
                actions.addWidget(button)
            body.addLayout(actions)
            card.setMinimumWidth(0)
            self._group_cards.append(card)
        self._layout_group_cards(self.width())

    def _layout_group_cards(self, width):
        while self.cards_grid.count():
            self.cards_grid.takeAt(0)
        columns = 3 if width >= 1000 else 2 if width >= 680 else 1
        for column in range(3):
            self.cards_grid.setColumnStretch(column, 1 if column < columns else 0)
        for index, card in enumerate(self._group_cards):
            self.cards_grid.addWidget(card, index // columns, index % columns)
        self.cards_empty.setVisible(not self._group_cards)
        if not self._group_cards:
            self.cards_grid.addWidget(self.cards_empty, 0, 0, 1, columns)

    def _actions_cell(self, edit_button, delete_button):
        cell = QWidget()
        row = QHBoxLayout(cell)
        row.setContentsMargins(2, 0, 2, 0)
        row.setSpacing(4)
        row.addWidget(edit_button, 0, Qt.AlignmentFlag.AlignCenter)
        row.addWidget(delete_button, 0, Qt.AlignmentFlag.AlignCenter)
        return cell

    @staticmethod
    def _template_cell(template):
        cell = QWidget()
        cell.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        cell.setStyleSheet(theme_qss('background: @surface; border: none;'))
        row = QHBoxLayout(cell)
        row.setContentsMargins(6, 0, 4, 0)
        row.setSpacing(6)
        marker = QLabel('\u2022')
        marker.setFixedWidth(8)
        marker.setStyleSheet(theme_qss('color: @accent; font-size: 12px; background: transparent; border: none;'))
        label = QLabel(str(template or ''))
        label.setToolTip(str(template or ''))
        label.setStyleSheet(theme_qss('color: @text; font-size: 11px; background: transparent; border: none;'))
        row.addWidget(marker)
        row.addWidget(label, 1)
        return cell

    def _load_linked_brands(self):
        self._linked_brands = None
        try:
            rows = self.db.conn.execute(
                'SELECT DISTINCT device_type, brand FROM device_brands '
                'WHERE COALESCE(is_active, 1)=1 ORDER BY id ASC'
            ).fetchall()
            self._linked_brands = {}
            for device_type, brand in rows:
                name = str(brand or '').strip()
                if name:
                    key = str(device_type or '').strip().casefold()
                    values = self._linked_brands.setdefault(key, [])
                    if name.casefold() not in {value.casefold() for value in values}:
                        values.append(name)
        except Exception:
            self._linked_brands = None

    def _linked_brands_for(self, name):
        if self._linked_brands is None:
            return None
        return self._linked_brands.get(name.strip().casefold(), [])

    def _refresh_metrics(self):
        columns = {row[1] for row in self.db.conn.execute('PRAGMA table_info(devices)')}
        device_where = 'WHERE COALESCE(is_deleted, 0)=0' if 'is_deleted' in columns else ''
        queries = (
            (self.metric_total_groups, "SELECT COUNT(*) FROM product_groups"),
            (self.metric_brand_pool, "SELECT COUNT(DISTINCT brand) FROM device_brands WHERE COALESCE(is_active, 1)=1"),
            (self.metric_device_inventory, "SELECT COUNT(*) FROM devices " + device_where),
        )
        for card, query in queries:
            try:
                self.db.cursor.execute(query)
                value = int(self.db.cursor.fetchone()[0] or 0)
            except Exception:
                value = None
            card.metric_value.setText('-' if value is None else f"{value:,}")
        try:
            active_service_count = self.db.conn.execute(
                'SELECT COUNT(*) FROM devices ' + device_where
                + " AND COALESCE(status, '') NOT IN ('Teslim Edildi', 'Iptal Edildi')"
                if device_where else
                "SELECT COUNT(*) FROM devices WHERE COALESCE(status, '') NOT IN ('Teslim Edildi', 'Iptal Edildi')"
            ).fetchone()[0]
            self.metric_device_inventory.metric_detail.setText(
                '{} aktif servis kay\u0131d\u0131'.format(active_service_count)
            )
        except Exception:
            self.metric_device_inventory.metric_detail.setText('Aktif servis kayd\u0131 yok')
        try:
            active_count = self.db.conn.execute(
                'SELECT COUNT(*) FROM product_groups WHERE COALESCE(is_active, 1)=1'
            ).fetchone()[0]
            recent_count = self.db.conn.execute(
                'SELECT COUNT(*) FROM product_groups WHERE created_at >= ?',
                ((datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d'),),
            ).fetchone()[0]
            self.active_group_summary.setText(
                '{} grup son 30 g\u00fcnde eklendi'.format(min(int(recent_count), int(active_count)))
            )
            self.active_category_badge.setText(f'{active_count} Aktif Kategori')
        except Exception:
            self.active_group_summary.setText('Aktif grup verisi kullanilamiyor')
            self.active_category_badge.setText('Aktif Kategori')
        try:
            row = self.db.conn.execute(
                'SELECT device_type, COUNT(*) FROM devices ' + device_where
                + ' GROUP BY device_type ORDER BY COUNT(*) DESC, device_type COLLATE NOCASE LIMIT 1'
            ).fetchone()
            if row and row[0]:
                self.metric_service_volume.metric_value.setText(str(row[0]))
                self.metric_service_volume.setToolTip('{} servis kaydi'.format(row[1]))
                total_services = int(self.db.conn.execute(
                    'SELECT COUNT(*) FROM devices ' + device_where
                ).fetchone()[0] or 0)
                share = round((int(row[1]) / total_services) * 100) if total_services else 0
                self.metric_service_volume.metric_detail.setText(
                    'Toplam hacmin %{}\'si'.format(share)
                )
            else:
                self.metric_service_volume.metric_value.setText('-')
                self.metric_service_volume.setToolTip('Servis kaydi yok')
                self.metric_service_volume.metric_detail.setText('Servis verisi yok')
        except Exception:
            self.metric_service_volume.metric_value.setText('-')
            self.metric_service_volume.setToolTip('Servis hacmi verisi kullanilamiyor')
            self.metric_service_volume.metric_detail.setText('Servis hacmi verisi yok')
        try:
            linked_groups = self.db.conn.execute(
                'SELECT COUNT(DISTINCT device_type) FROM device_brands WHERE COALESCE(is_active, 1)=1'
            ).fetchone()[0]
            self.metric_brand_pool.metric_detail.setText(
                '{} gruba ba\u011fl\u0131'.format(int(linked_groups or 0))
            )
        except Exception:
            self.metric_brand_pool.metric_detail.setText('Marka baglantisi yok')

    @classmethod
    def _standard_template(cls, group_name):
        normalized = str(group_name or '').casefold()
        for keywords, template in cls.STANDARD_TEMPLATES:
            if any(keyword in normalized for keyword in keywords):
                return template
        return 'Genel Servis \u015eablonu (F-001)'

    @classmethod
    def _group_description(cls, group_name):
        normalized = str(group_name or '').casefold()
        for keywords, description in cls.GROUP_DESCRIPTIONS:
            if any(keyword in normalized for keyword in keywords):
                return description
        return 'Cihaz ve servis kategorisi'

    @staticmethod
    def _brands_cell(summary):
        cell = QWidget()
        cell.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        cell.setStyleSheet(theme_qss('background: @surface; border: none;'))
        row = QHBoxLayout(cell)
        row.setContentsMargins(5, 0, 5, 0)
        label = QLabel(str(summary or '-'))
        label.setToolTip(str(summary or '-'))
        label.setStyleSheet(theme_qss(
            'QLabel { color: @text; background: @surface_alt; border: 1px solid @border; '
            'border-radius: 5px; padding: 3px 6px; font-size: 11px; }'
        ))
        row.addWidget(label, 0, Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
        row.addStretch(1)
        return cell

    def _group_name_cell(self, name):
        cell = QWidget()
        cell.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        cell.setStyleSheet(theme_qss('background: @surface; border: none;'))
        row = QHBoxLayout(cell)
        row.setContentsMargins(0, 4, 4, 4)
        row.setSpacing(9)
        normalized = str(name or '').casefold()
        icon_glyph, icon_color = ('\U0001f4e6', '@accent')
        for keywords, glyph, color in (
            (('laptop', 'bilgisayar', 'pc'), '\U0001f4bb', '@finance_primary'),
            (('elektronik',), '\u2699', '@finance_action_income'),
            (('telefon', 'mobil'), '\U0001f4f1', '@accent'),
            (('all in one', 'aio'), '\U0001f5a5', '@finance_action_income'),
            (('saat', 'giyilebilir'), '\U0001f4f1', '@finance_action_expense'),
            (('priz', 'iot'), '\U0001f50c', '@warning'),
            (('kamera', 'cctv'), '\U0001f4bb', '@accent'),
            (('merkez', 'hub'), '\U0001f3e0', '@accent'),
            (('ayd\u0131nlatma', 'led'), '\U0001f4ca', '@warning'),
        ):
            if any(keyword in normalized for keyword in keywords):
                icon_glyph, icon_color = glyph, color
                break
        icon = NavigationGlyph(
            icon_glyph, title=str(name or ''), active_token=icon_color.lstrip('@')
        )
        icon.setFixedSize(28, 28)
        icon.set_navigation_state(True)
        icon.setStyleSheet(theme_qss(
            'QLabel { color: %s; background: @surface_alt; border: 1px solid %s; '
            'border-radius: 7px; font-size: 13px; }' % (icon_color, icon_color)
        ))
        stack = QVBoxLayout()
        stack.setContentsMargins(0, 0, 0, 0)
        stack.setSpacing(1)
        title = QLabel(str(name or ''))
        title.setStyleSheet(theme_qss('color: @text; background: transparent; border: none; font-weight: 700;'))
        subtitle = QLabel(self._group_description(name))
        subtitle.setStyleSheet(theme_qss('color: @text_muted; background: transparent; border: none; font-size: 10px;'))
        stack.addWidget(title)
        stack.addWidget(subtitle)
        row.addWidget(icon)
        row.addLayout(stack, 1)
        return cell

    def show_group_hierarchy(self):
        dialog = ModernDialog(title='Grup Hiyerar\u015fisi', parent=self, width=680, height=560)
        tree = QTreeWidget()
        tree.setHeaderHidden(True)
        tree.setRootIsDecorated(True)
        tree.setStyleSheet(theme_qss(
            'QTreeWidget { background: @surface; color: @text; border: 1px solid @border; '
            'border-radius: 10px; padding: 8px; }'
        ))
        for index in range(self.table.rowCount()):
            name_item = self.table.item(index, 1)
            if not name_item:
                continue
            group = QTreeWidgetItem([name_item.text()])
            tree.addTopLevelItem(group)
            for brand in self._linked_brands_for(name_item.text()) or []:
                QTreeWidgetItem(group, [brand])
            group.setExpanded(True)
        dialog.content_layout.addWidget(tree)
        dialog.add_button('Kapat', variant='secondary', callback=dialog.accept)
        dialog.exec()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._apply_responsive_layout()
        self._fit_table_card()
        if hasattr(self, 'cards_grid'):
            self._layout_group_cards(event.size().width())

    def showEvent(self, event):
        super().showEvent(event)
        self._apply_responsive_layout()

    def _apply_responsive_layout(self):
        window = self.main_window or self.window()
        sidebar = getattr(window, 'app_sidebar', None)
        page_width = window.width() - (sidebar.width() if sidebar else 192)
        side_margin = max(16, min(24, (page_width - 924) // 2))
        reference_page = getattr(window, "_current_page_index", None) == 146
        if self.width() <= 1200:
            self._root_layout.setContentsMargins(side_margin + 12, 21, side_margin, 12)
            self._root_layout.setSpacing(14)
        else:
            self._root_layout.setContentsMargins(
                side_margin, 19 if reference_page else 24,
                side_margin, 12 if reference_page else 20,
            )

    def add_group(self):
        name, accepted = ModernInputDialog.get_text(
            self,
            "Yeni \u00dcr\u00fcn Grubu",
            "\u00dcr\u00fcn grubu ad\u0131:",
        )
        if not accepted:
            return
        name = str(name or "").strip()
        if not name:
            show_warning(self, "Grup ad\u0131 bo\u015f olamaz.")
            return
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        try:
            self.db.cursor.execute(
                """
                INSERT INTO product_groups (
                    name, is_active, created_at, updated_at
                )
                VALUES (?, 1, ?, ?)
                """,
                (name, now, now),
            )
            self.db.conn.commit()
            self.load_data()
            show_success(self, "\u00dcr\u00fcn grubu eklendi.")
        except Exception as exc:
            show_error(self, f"Kay\u0131t eklenemedi: {exc}")

    def edit_group(self, group_id, current_name):
        name, accepted = ModernInputDialog.get_text(
            self,
            "\u00dcr\u00fcn Grubunu D\u00fczenle",
            "\u00dcr\u00fcn grubu ad\u0131:",
            current_name,
        )
        if not accepted:
            return
        name = str(name or "").strip()
        if not name:
            show_warning(self, "Grup ad\u0131 bo\u015f olamaz.")
            return
        try:
            self.db.cursor.execute(
                "UPDATE product_groups SET name=?, updated_at=? WHERE id=?",
                (
                    name,
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    group_id,
                ),
            )
            self.db.cursor.execute(
                """
                UPDATE device_brands
                SET device_type=?
                WHERE device_type=?
                """,
                (name, current_name),
            )
            self.db.conn.commit()
            self.load_data()
            show_info(self, "\u00dcr\u00fcn grubu g\u00fcncellendi.")
        except Exception as exc:
            self.db.conn.rollback()
            show_error(self, f"Kay\u0131t g\u00fcncellenemedi: {exc}")

    def delete_group(self, group_id, name):
        dialog = SimpleConfirmDialog(
            self,
            title="\u00dcr\u00fcn Grubunu Sil",
            text=f"'{name}' \u00fcr\u00fcn grubu silinsin mi?",
        )
        if not dialog.exec():
            return
        try:
            self.db.cursor.execute(
                "DELETE FROM product_groups WHERE id=?",
                (group_id,),
            )
            self.db.conn.commit()
            self.load_data()
            show_success(self, "\u00dcr\u00fcn grubu silindi.")
        except Exception as exc:
            self.db.conn.rollback()
            show_error(self, f"Kay\u0131t silinemedi: {exc}")

    def previous_page(self):
        if self.current_page > 0:
            self.current_page -= 1
            self.load_data()

    def next_page(self):
        if (self.current_page + 1) * self.PAGE_SIZE < self.total_count:
            self.current_page += 1
            self.load_data()

    def _all_names(self):
        where_sql, params = self._where_clause()
        self.db.cursor.execute(
            f"""
            SELECT name, is_active
            FROM product_groups
            {where_sql}
            ORDER BY name COLLATE NOCASE
            """,
            tuple(params),
        )
        return [
            (str(row[0] or ""), "Aktif" if row[1] else "Pasif")
            for row in (self.db.cursor.fetchall() or [])
        ]

    def _export_rows(self):
        rows = []
        for name, status in self._all_names():
            brands = self._linked_brands_for(name)
            rows.append((name, '-' if brands is None else ', '.join(brands), status))
        return rows

    def export_excel(self):
        path, _selected = QFileDialog.getSaveFileName(
            self,
            "Excel \u00c7\u0131kt\u0131",
            "urun_gruplari.xlsx",
            "Excel (*.xlsx)",
        )
        if not path:
            return
        try:
            from openpyxl import Workbook

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Urun Gruplari"
            sheet.append(["Urun Grubu", "Bagli Marka / Kategori", "Durum"])
            for row in self._export_rows():
                sheet.append(list(row))
            workbook.save(path)
            show_success(self, f"Excel kaydedildi: {path}")
        except Exception as exc:
            show_error(self, f"Excel olusturulamadi: {exc}")

    def _html_document(self):
        rows = "".join(
            "<tr><td>{}</td><td>{}</td><td>{}</td></tr>".format(
                html.escape(name),
                html.escape(brands),
                html.escape(status),
            )
            for name, brands, status in self._export_rows()
        )
        return (
            "<h2>Urun Grubu Yonetimi</h2>"
            "<table border='1' cellspacing='0' cellpadding='6' "
            "width='100%'><tr><th>Urun Grubu</th><th>Bagli Marka / Kategori</th><th>Durum</th></tr>"
            f"{rows}</table>"
        )

    def export_pdf(self):
        path, _selected = QFileDialog.getSaveFileName(
            self,
            "PDF Kaydet",
            "urun_gruplari.pdf",
            "PDF (*.pdf)",
        )
        if not path:
            return
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(str(Path(path)))
        document = QTextDocument()
        document.setHtml(self._html_document())
        document.print(printer)
        show_success(self, f"PDF kaydedildi: {path}")

    def print_list(self):
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        dialog = QPrintDialog(printer, self)
        if not dialog.exec():
            return
        document = QTextDocument()
        document.setHtml(self._html_document())
        document.print(printer)
